"""Historical data loader: load a folder of historical telemetry files for one site / battery system through the
bulk-load APIs. Generic - works for any site whose master data is loaded.

Inputs
  --folder      folder with the files (below <project>\\input); sub-folders are scanned too unless --no-recursive
  --system      battery system name (battery_system.system_name), e.g. EESL-BESS-S02
  --ems-code / --hvac-code   the EMS controller / HVAC unit the EMS and HVAC files belong to (their files carry
                no asset column); default: the system's only EMS / HVAC
  --feeds       which feeds to load (default all): bms, pcs, ems, hvac, module, alarm, dispatch
  --files-json  JSON file with the list of files to load (names relative to the folder, as the scan shows them);
                every other file of the chosen feeds is left out (EXCLUDED). Default: all files that pass the check.
  --progress    JSON file the run keeps updated (the page reads it when the user refreshes)

Files are assigned to feeds by name (e.g. *BMS*, *PCS*, *Module*, *Alarm*, *Dispatch*; .xlsx / .xlsm / .csv).
Before anything is loaded every file is checked against the system's master data: the BMS system name, the PCS
codes, module codes and the dispatch EMS codes in the file must belong to the chosen system (a file for another
system is skipped, never loaded). Feeds are then loaded in order (BMS, PCS, EMS, HVAC, modules, alarms, dispatch
events), files of a feed in name order. Every bulk API inserts new rows and updates existing ones (same asset +
timestamp / same id), so a folder can be loaded again safely.

One battery_raw.batchlog row is written per run (job_name 'historical_load') with the result of every file.
"""
import argparse
import json
import os
import re
import socket
import sys
import time
import urllib.error
import urllib.request
import uuid
from dataclasses import dataclass, field
from itertools import islice
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Optional

# the bulk APIs are deny-by-default: authenticate as the batch service account (see bess_service_auth.py)
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from bess_service_auth import auth_headers  # noqa: E402
from bess_common import observability as obs  # noqa: E402  (JSON logs, operations, activity log)

from urllib.parse import unquote, urlparse

import psycopg2
from psycopg2.extras import Json

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
INPUT_ROOT = PROJECT / "input"
DEFAULT_API = "http://127.0.0.1:8000"  # not localhost: ~2 s per call (IPv6 first)
JOB_NAME = "historical_load"
LOCK = HERE / ".historical_load.lock"
STALE_LOCK_S = 6 * 3600
PLANT_TZ = "Asia/Kolkata"
EXTENSIONS = (".xlsx", ".xlsm", ".csv")


@dataclass(frozen=True)
class Feed:
    key: str
    label: str
    endpoint: str
    pattern: str  # regex on the file name (case-insensitive)
    code_column: Optional[str] = None  # column checked against the system's master data
    form: Optional[str] = None  # form field the endpoint needs (ems_code / hvac_code)


FEEDS: tuple[Feed, ...] = (
    Feed("bms", "BMS", "/api/system-telemetry/bulk", r"bms|system[_ -]?telemetry", "system_name"),
    Feed("pcs", "PCS", "/api/pcs-telemetry/bulk", r"pcs", "pcs_code"),
    Feed("ems", "EMS", "/api/ems-telemetry/bulk", r"ems(?!.*dispatch)", form="ems_code"),
    Feed("hvac", "HVAC", "/api/hvac-telemetry/bulk", r"hvac", form="hvac_code"),
    Feed("module", "Module", "/api/module-telemetry/bulk", r"module", "module_code"),
    Feed("alarm", "Alarms", "/api/alarm-events/bulk", r"alarm"),
    Feed("dispatch", "Dispatch events", "/api/dispatch-events/bulk", r"dispatch", "ems_code"),
)
FEED_KEYS = tuple(f.key for f in FEEDS)


class LoadError(Exception):
    pass


# ---------------------------------------------------------------- setup


def database_dsn() -> dict:
    url = None
    for line in (PROJECT / "backend" / ".env").read_text(encoding="utf-8").splitlines():
        if line.strip().startswith("DATABASE_URL="):
            url = line.split("=", 1)[1].strip()
    if not url:
        raise LoadError("DATABASE_URL not found in backend/.env")
    u = urlparse(url.replace("postgresql+psycopg2://", "postgresql://"))
    return dict(host=u.hostname, port=u.port or 5432, dbname=u.path.lstrip("/"), user=unquote(u.username or ""),
                password=unquote(u.password or ""), options=f"-c timezone={PLANT_TZ}")


def resolve_folder(folder: str) -> Path:
    """A folder inside <project>\\input (absolute, or relative to it)."""
    path = Path(folder)
    path = (path if path.is_absolute() else INPUT_ROOT / path).resolve()
    if INPUT_ROOT.resolve() not in (path, *path.parents):
        raise LoadError(f"The folder must be inside {INPUT_ROOT}")
    if not path.is_dir():
        raise LoadError(f"Folder not found: {path}")
    return path


# ---------------------------------------------------------------- master data of the system


def system_assets(conn, system_name: str) -> dict:
    with conn.cursor() as cur:
        cur.execute("""SELECT b.battery_system_id, b.system_name, s.site_name FROM battery_master.battery_system b
                       JOIN battery_master.site s USING (site_id) WHERE b.system_name = %s""", (system_name,))
        row = cur.fetchone()
        if not row:
            raise LoadError(f"Battery system '{system_name}' not found")
        system_id = row[0]
        q = lambda sql: [r[0] for r in (cur.execute(sql, (system_id,)) or cur.fetchall())]  # noqa: E731
        return {
            "system_id": system_id, "system_name": row[1], "site_name": row[2],
            "pcs": q("SELECT pcs_code FROM battery_master.pcs WHERE system_id = %s ORDER BY 1"),
            "ems": q("SELECT ems_code FROM battery_master.ems_controller WHERE battery_system_id = %s ORDER BY 1"),
            "hvac": q("SELECT hvac_code FROM battery_master.hvac_unit WHERE battery_system_id = %s ORDER BY 1"),
            "modules": q("""SELECT m.serial_number FROM battery_master.battery_module m
                            JOIN battery_master.battery_pack p USING (pack_id)
                            JOIN battery_master.battery_container c USING (container_id)
                            WHERE c.battery_system_id = %s AND m.serial_number IS NOT NULL"""),
        }


# ---------------------------------------------------------------- scan


def feed_of(name: str) -> Optional[Feed]:
    low = name.lower()
    for feed in FEEDS:
        if feed.key == "ems" and ("dispatch" in low or "hvac" in low):
            continue
        if re.search(feed.pattern, low):
            return feed
    return None


def read_column(path: Path, column: str, limit: int = 2_000) -> Optional[set[str]]:
    """Distinct values of one column in the first `limit` rows (a sample is enough to tell which system a file is
    for); None when the file lacks the column."""
    if path.suffix.lower() == ".csv":
        import csv
        with path.open(encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            header = [h.strip().lower() for h in next(reader, [])]
            if column not in header:
                return None
            i = header.index(column)
            return {row[i].strip() for row in islice(reader, limit) if len(row) > i and row[i].strip()}
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = next((wb[n] for n in wb.sheetnames if n.strip().lower() in ("data", "dispatch events")), wb.worksheets[0])
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows, [])]
    if column not in header:
        wb.close()
        return None
    i = header.index(column)
    values = {str(r[i]).strip() for r in islice(rows, limit) if i < len(r) and r[i] not in (None, "")}
    wb.close()
    return values


def scan(folder: Path, assets: dict, feeds: tuple[str, ...] = FEED_KEYS, recursive: bool = True,
         ems_code: Optional[str] = None, hvac_code: Optional[str] = None) -> dict:
    """Files per feed with their check result: {feeds: [{key, label, files: [{path, name, size, ok, message}]}],
    ignored: [...], problems: [...]}."""
    files = sorted(p for p in (folder.rglob("*") if recursive else folder.glob("*"))
                   if p.is_file() and p.suffix.lower() in EXTENSIONS and not p.name.startswith(("~$", ".")))
    grouped: dict[str, list[Path]] = {k: [] for k in FEED_KEYS}
    ignored = []
    for p in files:
        feed = feed_of(p.name)
        if feed is None:
            ignored.append(str(p.relative_to(folder)))
        else:
            grouped[feed.key].append(p)
    problems = []
    ems_code = ems_code or (assets["ems"][0] if len(assets["ems"]) == 1 else None)
    hvac_code = hvac_code or (assets["hvac"][0] if len(assets["hvac"]) == 1 else None)
    if grouped["ems"] and "ems" in feeds and ems_code not in assets["ems"]:
        problems.append(f"Choose the EMS controller for the EMS files (system has: {', '.join(assets['ems']) or 'none'})")
    if grouped["hvac"] and "hvac" in feeds and hvac_code not in assets["hvac"]:
        problems.append(f"Choose the HVAC unit for the HVAC files (system has: {', '.join(assets['hvac']) or 'none'})")

    expected = {"bms": {assets["system_name"]}, "pcs": set(assets["pcs"]), "module": set(assets["modules"]),
                "dispatch": set(assets["ems"])}
    out = []
    for feed in FEEDS:
        entries = []
        for p in grouped[feed.key]:
            entry = {"path": str(p), "name": str(p.relative_to(folder)), "size": p.stat().st_size, "ok": True,
                     "message": None, "selected": feed.key in feeds}
            if feed.code_column and feed.key in feeds:
                try:
                    values = read_column(p, feed.code_column)
                except Exception as exc:  # noqa: BLE001 - reported per file
                    values, entry["ok"], entry["message"] = None, False, f"Can't read the file: {exc}"
                if values is not None:
                    foreign = sorted(values - expected[feed.key])
                    if foreign and feed.key == "module" and len(foreign) < len(values):
                        entry["message"] = f"{len(foreign)} module code(s) not in this system - those rows will fail"
                    elif foreign:
                        entry["ok"] = False
                        entry["message"] = (f"{feed.code_column} {', '.join(foreign[:3])}{' ...' if len(foreign) > 3 else ''}"
                                            f" is not part of {assets['system_name']} - file skipped")
                elif entry["ok"]:
                    entry["ok"], entry["message"] = False, f"No {feed.code_column} column"
            if feed.form == "ems_code":
                entry["message"] = entry["message"] or f"loaded for EMS {ems_code}"
            if feed.form == "hvac_code":
                entry["message"] = entry["message"] or f"loaded for HVAC {hvac_code}"
            entries.append(entry)
        out.append({"key": feed.key, "label": feed.label, "endpoint": feed.endpoint, "files": entries})
    return {"folder": str(folder), "feeds": out, "ignored": ignored, "problems": problems,
            "ems_code": ems_code, "hvac_code": hvac_code}


# ---------------------------------------------------------------- load


def post_file(api: str, endpoint: str, path: Path, form: dict[str, str]) -> tuple[int, dict]:
    boundary = uuid.uuid4().hex
    parts = [f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode() for k, v in form.items()]
    ctype = "text/csv" if path.suffix.lower() == ".csv" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{path.name}"\r\n'
                 f"Content-Type: {ctype}\r\n\r\n".encode() + path.read_bytes() + b"\r\n")
    parts.append(f"--{boundary}--\r\n".encode())
    request = urllib.request.Request(api + endpoint, data=b"".join(parts), method="POST",
                                     headers={"Content-Type": f"multipart/form-data; boundary={boundary}", **auth_headers(api)})
    try:
        with urllib.request.urlopen(request, timeout=1800) as response:
            return response.status, json.loads(response.read() or b"{}")
    except urllib.error.HTTPError as exc:
        try:
            return exc.code, json.loads(exc.read() or b"{}")
        except ValueError:
            return exc.code, {"detail": exc.reason}
    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        return 0, {"detail": f"API not reachable: {exc}"}


@dataclass
class Progress:
    path: Optional[Path]
    state: dict = field(default_factory=dict)

    def write(self, **changes) -> None:
        self.state.update(changes, updated_at=datetime.now().astimezone().isoformat(timespec="seconds"))
        if not self.path:
            return
        # the backend may be reading the file at this moment (Windows then refuses the replace): retry, and never let
        # a progress update stop the load
        tmp = self.path.with_suffix(".tmp")
        for attempt in range(20):
            try:
                tmp.write_text(json.dumps(self.state, indent=2, default=str), encoding="utf-8")
                os.replace(tmp, self.path)
                return
            except OSError:
                time.sleep(0.05 * (attempt + 1))


def run(folder: str, system: str, feeds: tuple[str, ...] = FEED_KEYS, recursive: bool = True,
        ems_code: Optional[str] = None, hvac_code: Optional[str] = None, api: str = DEFAULT_API,
        progress_path: Optional[Path] = None, log: Callable[[str], None] = print,
        include: Optional[set[str]] = None) -> dict:
    """include: file names (relative to the folder) to load; None = every file that passes the check."""
    run_id = str(uuid.uuid4())
    started = datetime.now().astimezone()
    progress = Progress(progress_path, {"run_id": run_id, "status": "RUNNING", "started_at": started.isoformat(),
                                        "folder": folder, "system": system, "feeds": list(feeds), "files": [],
                                        "current": None, "error": None, "batch_id": None})
    progress.write()
    dsn = database_dsn()
    report: dict[str, Any] = {"run_id": run_id, "folder": folder, "system": system, "feeds": list(feeds), "files": [],
                              "warnings": [], "error": None}
    try:
        path = resolve_folder(folder)
        with psycopg2.connect(**dsn) as conn:
            assets = system_assets(conn, system)
        conn.close()
        plan = scan(path, assets, feeds, recursive, ems_code, hvac_code)
        report.update(site=assets["site_name"], ems_code=plan["ems_code"], hvac_code=plan["hvac_code"],
                      ignored=plan["ignored"])
        if plan["problems"]:
            raise LoadError("; ".join(plan["problems"]))
        todo = [(f, e) for f in FEEDS for g in plan["feeds"] if g["key"] == f.key and f.key in feeds for e in g["files"]]
        if not todo:
            raise LoadError("No files for the chosen feeds in this folder")
        if include is not None:
            unknown = sorted(include - {e["name"] for _, e in todo})
            if unknown:
                report["warnings"].append(f"Selected file(s) not found for the chosen feeds: {', '.join(unknown[:5])}")
            if not any(e["name"] in include for _, e in todo):
                raise LoadError("None of the selected files is in the folder any more")

        def initial(e: dict) -> tuple[str, Optional[str]]:
            if include is not None and e["name"] not in include:
                return "EXCLUDED", "left out by the user"
            return ("PENDING", None) if e["ok"] else ("SKIPPED", e["message"])

        files = [{"feed": f.label, "file": e["name"], "status": initial(e)[0], "message": initial(e)[1]} for f, e in todo]
        progress.write(files=files, site=assets["site_name"])
        log(f"{assets['site_name']} / {system}: {len(todo)} file(s) in {path}")
        for i, (feed, entry) in enumerate(todo):
            item = files[i]
            if item["status"] == "EXCLUDED":
                report["files"].append({**item, "rows": 0, "inserted": 0, "updated": 0, "skipped": 0})
                continue
            if not entry["ok"]:
                log(f"  skip {entry['name']}: {entry['message']}")
                report["files"].append({**item, "rows": 0, "inserted": 0, "updated": 0, "skipped": 0})
                continue
            item["status"] = "LOADING"
            progress.write(files=files, current=entry["name"])
            form = {}
            if feed.form == "ems_code":
                form["ems_code"] = plan["ems_code"]
            if feed.form == "hvac_code":
                form["hvac_code"] = plan["hvac_code"]
            t0 = time.perf_counter()
            status, body = post_file(api, feed.endpoint, Path(entry["path"]), form)
            seconds = round(time.perf_counter() - t0, 1)
            if status == 200:
                total = body.get("total_rows", 0)
                item.update(status="LOADED" if not body.get("error_count") else "PARTIAL", http=status, rows=total,
                            inserted=body.get("inserted", 0), updated=body.get("updated", 0),
                            skipped=total - body.get("inserted", 0) - body.get("updated", 0), seconds=seconds,
                            errors=[f"row {e.get('row')}: {'; '.join(e.get('errors', []))}" for e in body.get("errors", [])][:20])
            else:
                detail = body.get("detail") if isinstance(body, dict) else body
                item.update(status="FAILED", http=status, rows=0, inserted=0, updated=0, skipped=0, seconds=seconds,
                            message=str(detail)[:500])
                if status == 0:
                    report["error"] = str(detail)
            log(f"  {feed.label:15} {entry['name']}: {item['status']} {item.get('inserted', 0)} ins / "
                f"{item.get('updated', 0)} upd / {item.get('skipped', 0)} skipped ({seconds}s)")
            report["files"].append(dict(item))
            progress.write(files=files, current=None)
            if status == 0:  # the API is down: stop, the rest would fail too
                for rest in files[i + 1:]:
                    if rest["status"] != "EXCLUDED":
                        rest.update(status="SKIPPED", message="not attempted: API not reachable")
                    report["files"].append({**rest, "rows": 0, "inserted": 0, "updated": 0, "skipped": 0})
                break
        # daily summaries: refresh the aggregates of the feeds that loaded rows
        loaded_feeds = sorted({f.key for f, _ in todo for r in report["files"]
                               if r["feed"] == f.label and r["status"] in ("LOADED", "PARTIAL")})
        if loaded_feeds:
            progress.write(current="Refreshing daily summaries")
            try:
                report["aggregates_refreshed"] = refresh_aggregates(dsn, assets["system_id"], loaded_feeds, log)
            except psycopg2.Error as exc:
                report["warnings"].append(f"Daily summaries not refreshed: {exc}")
    except LoadError as exc:
        report["error"] = str(exc)
        log(f"error: {exc}")

    finished = datetime.now().astimezone()
    excluded = sum(1 for f in report["files"] if f["status"] == "EXCLUDED")
    report["files"] = [f for f in report["files"] if f["status"] != "EXCLUDED"] +         [f for f in report["files"] if f["status"] == "EXCLUDED"]
    report["excluded"] = excluded
    loaded = [f for f in report["files"] if f["status"] in ("LOADED", "PARTIAL")]
    bad = [f for f in report["files"] if f["status"] in ("FAILED", "SKIPPED", "PARTIAL")]
    status = "FAILED" if report["error"] and not loaded or not loaded else ("SUCCESS" if not bad else "PARTIAL")
    if status != "SUCCESS" and not report["error"]:
        report["error"] = f"{len(bad)} of {len(report['files'])} file(s) not fully loaded"
    totals = dict(
        feeds_total=len(report["files"]) - excluded, feeds_loaded=len(loaded),
        feeds_failed=sum(1 for f in report["files"] if f["status"] in ("FAILED", "SKIPPED")),
        rows_total=sum(f.get("rows", 0) for f in report["files"]),
        rows_inserted=sum(f.get("inserted", 0) for f in report["files"]),
        rows_updated=sum(f.get("updated", 0) for f in report["files"]),
        rows_skipped=sum(max(f.get("skipped", 0), 0) for f in report["files"]),
    )
    feed_results = [{"feed": f["feed"], "csv": f["file"], "rows": f.get("rows", 0), "http": f.get("http"),
                     "inserted": f.get("inserted"), "updated": f.get("updated"), "skipped": f.get("skipped"),
                     "result": "Excluded by the user" if f["status"] == "EXCLUDED" else None,
                     "errors": [] if f["status"] == "EXCLUDED" else
                     f.get("errors") or ([f["message"]] if f.get("message") and f["status"] != "LOADED" else [])}
                    for f in report["files"]]
    batch_id = None
    try:
        with psycopg2.connect(**dsn) as conn, conn.cursor() as cur:
            cur.execute(
                """INSERT INTO battery_raw.batchlog
                     (run_id, job_name, file_name, status, started_at, finished_at, duration_ms, feeds_total,
                      feeds_loaded, feeds_failed, rows_total, rows_inserted, rows_updated, rows_skipped, feed_results,
                      warnings, id_changes, error_message, archive_path, api_url, host_name, report)
                   VALUES (%(run_id)s, %(job)s, %(file_name)s, %(status)s, %(started)s, %(finished)s, %(ms)s,
                      %(feeds_total)s, %(feeds_loaded)s, %(feeds_failed)s, %(rows_total)s, %(rows_inserted)s,
                      %(rows_updated)s, %(rows_skipped)s, %(feed_results)s, %(warnings)s, '[]'::jsonb, %(error)s,
                      %(folder)s, %(api)s, %(host)s, %(report)s) RETURNING batch_id""",
                dict(run_id=run_id, job=JOB_NAME, file_name=f"{report.get('site') or system} · {Path(folder).name}",
                     status=status, started=started, finished=finished,
                     ms=int((finished - started).total_seconds() * 1000), **totals, feed_results=Json(feed_results),
                     warnings=Json(report["warnings"]), error=report["error"], folder=folder, api=api,
                     host=socket.gethostname(), report=Json(json.loads(json.dumps(report, default=str)))))
            batch_id = cur.fetchone()[0]
        conn.close()
    except psycopg2.Error as exc:
        log(f"batch log not written: {exc}")
    progress.write(status=status, finished_at=finished.isoformat(), error=report["error"], batch_id=batch_id,
                   totals=totals, current=None)
    log(f"{status}: {totals['rows_inserted']} inserted, {totals['rows_updated']} updated, "
        f"{totals['rows_skipped']} skipped | batchlog #{batch_id}")
    return {"status": status, "batch_id": batch_id, **totals}


# daily continuous aggregates (TimescaleDB) per feed, and the SQL giving the system's time range in that feed.
# Their scheduled refresh only covers the last 3 days, so a historical load must refresh them for its range.
AGGREGATES = {
    "bms": ("system_telemetry_daily", "SELECT min(ts), max(ts) FROM battery_telemetry.system_telemetry WHERE system_id = %s"),
    "pcs": ("pcs_telemetry_daily", """SELECT min(t.ts), max(t.ts) FROM battery_telemetry.pcs_telemetry t
                                      JOIN battery_master.pcs p USING (pcs_id) WHERE p.system_id = %s"""),
    "ems": ("ems_telemetry_daily", """SELECT min(t.ts), max(t.ts) FROM battery_telemetry.ems_telemetry t
                                      JOIN battery_master.ems_controller e USING (ems_id) WHERE e.battery_system_id = %s"""),
    "hvac": ("hvac_telemetry_daily", """SELECT min(t.ts), max(t.ts) FROM battery_telemetry.hvac_telemetry t
                                        JOIN battery_master.hvac_unit h USING (hvac_id) WHERE h.battery_system_id = %s"""),
    "module": ("module_telemetry_daily", """SELECT min(t.ts), max(t.ts) FROM battery_telemetry.module_telemetry t
                                            JOIN battery_master.battery_module m USING (module_id)
                                            JOIN battery_master.battery_pack p USING (pack_id)
                                            JOIN battery_master.battery_container c USING (container_id)
                                            WHERE c.battery_system_id = %s"""),
}


def refresh_aggregates(dsn: dict, system_id: int, feeds: list[str], log: Callable[[str], None] = print) -> list[dict]:
    """Refresh the daily aggregates of the loaded feeds over the system's whole time range (plus a day each side)."""
    done = []
    conn = psycopg2.connect(**dsn)
    conn.autocommit = True  # refresh_continuous_aggregate can't run inside a transaction
    try:
        with conn.cursor() as cur:
            for key in feeds:
                if key not in AGGREGATES:
                    continue
                view, range_sql = AGGREGATES[key]
                cur.execute(range_sql, (system_id,))
                lo, hi = cur.fetchone()
                if lo is None:
                    continue
                t0 = time.perf_counter()
                cur.execute("CALL refresh_continuous_aggregate(%s, %s::timestamptz - interval '1 day', "
                            "%s::timestamptz + interval '1 day')", (f"battery_telemetry.{view}", lo, hi))
                seconds = round(time.perf_counter() - t0, 1)
                done.append({"view": view, "from": lo.isoformat(), "to": hi.isoformat(), "seconds": seconds})
                log(f"  refreshed {view} {lo:%Y-%m-%d} .. {hi:%Y-%m-%d} ({seconds}s)")
    finally:
        conn.close()
    return done


class Lock:
    def __enter__(self) -> bool:
        if LOCK.exists() and time.time() - LOCK.stat().st_mtime > STALE_LOCK_S:
            LOCK.unlink(missing_ok=True)
        try:
            fd = os.open(LOCK, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError:
            self.held = False
            return False
        with os.fdopen(fd, "w") as f:
            f.write(f"pid {os.getpid()} since {datetime.now():%Y-%m-%d %H:%M:%S}\n")
        self.held = True
        return True

    def __exit__(self, *exc) -> None:
        if self.held:
            LOCK.unlink(missing_ok=True)


def main(argv: Optional[list[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="Load a folder of historical telemetry files for one battery system.")
    ap.add_argument("--folder", required=True, help="folder inside <project>\\input, e.g. \"historical\\Bhadla BESS\"")
    ap.add_argument("--system", required=True, help="battery system name, e.g. EESL-BESS-S02")
    ap.add_argument("--feeds", default=",".join(FEED_KEYS), help=f"comma list of {', '.join(FEED_KEYS)}")
    ap.add_argument("--ems-code")
    ap.add_argument("--hvac-code")
    ap.add_argument("--no-recursive", action="store_true")
    ap.add_argument("--api", default=DEFAULT_API)
    ap.add_argument("--progress", type=Path)
    ap.add_argument("--files-json", type=Path, help="JSON list of the files to load (relative to the folder)")
    args = ap.parse_args(argv)
    feeds = tuple(f.strip() for f in args.feeds.split(",") if f.strip())
    unknown = set(feeds) - set(FEED_KEYS)
    if unknown:
        ap.error(f"unknown feed(s): {', '.join(sorted(unknown))}")
    with Lock() as locked:
        if not locked:
            print("another historical load is still running")
            return 4
        include = set(json.loads(args.files_json.read_text(encoding="utf-8"))) if args.files_json else None
        obs.setup_logging("loader-historical")
        obs.enable_activity_log(obs.engine_from_dsn(database_dsn()))
        with obs.request_context(str(uuid.uuid4()), user="batch.service@bess.local", role="BATCH_SERVICE"),                 obs.operation("load", "historical_load", system=args.system, folder=args.folder, feeds=list(feeds)) as op:
            result = run(args.folder, args.system, feeds, not args.no_recursive, args.ems_code, args.hvac_code, args.api,
                         args.progress, log=lambda m: print(m, flush=True), include=include)
            op.set(rows_total=result.get("rows_total", 0), rows_inserted=result.get("rows_inserted", 0),
                   rows_updated=result.get("rows_updated", 0), rows_rejected=result.get("rows_skipped", 0),
                   status=result["status"], batch_id=result.get("batch_id"),
                   outcome="success" if result["status"] == "SUCCESS" else "failed")
        obs.activity.flush()
    return {"SUCCESS": 0, "PARTIAL": 1}.get(result["status"], 2)


if __name__ == "__main__":
    sys.exit(main())
