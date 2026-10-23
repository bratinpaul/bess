import csv
import io
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Optional

import openpyxl
from dateutil import parser as date_parser
from pydantic import ValidationError
from sqlalchemy import literal_column
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

from bess_common.errors import FileFormatError
from app import models, schemas

MAX_ROWS = 200_000

TIMESTAMP_HEADERS = ("ts", "timestamp")

# Maps a normalized (lowercased) source column header to an ems_telemetry field.
COLUMN_MAP = {
    "dispatch_mode": "dispatch_mode",
    "grid_service": "grid_service",
    "market_price_mwh": "market_price_mwh",
    "p_setpoint_kw": "p_setpoint_kw",
    "poi_active_power_kw": "poi_active_power_kw",
    "poi_reactive_power_kvar": "poi_reactive_power_kvar",
    "tracking_error_kw": "tracking_error_kw",
    "grid_frequency_hz": "grid_frequency_hz",
    "poi_voltage_kv": "poi_voltage_kv",
    "agc_signal": "agc_signal",
    "plant_soc_pct": "plant_soc_pct",
    "target_soc_pct": "target_soc_pct",
    "avail_charge_kw": "avail_charge_kw",
    "avail_discharge_kw": "avail_discharge_kw",
    "avail_charge_energy_kwh": "avail_charge_energy_kwh",
    "avail_discharge_energy_kwh": "avail_discharge_energy_kwh",
    "aux_load_kw": "aux_load_kw",
    "aux_load_pct": "aux_load_pct",
    "cum_import_kwh": "cum_import_kwh",
    "cum_export_kwh": "cum_export_kwh",
    "revenue_cum_usd": "revenue_cum_usd",
    "units_online": "units_online",
    "plant_available": "plant_available",
    "alarm_level": "alarm_level",
}

INT_FIELDS = {"units_online", "plant_available", "alarm_level"}

DECIMAL_FIELDS = {
    "market_price_mwh",
    "p_setpoint_kw",
    "poi_active_power_kw",
    "poi_reactive_power_kvar",
    "tracking_error_kw",
    "grid_frequency_hz",
    "poi_voltage_kv",
    "agc_signal",
    "plant_soc_pct",
    "target_soc_pct",
    "avail_charge_kw",
    "avail_discharge_kw",
    "avail_charge_energy_kwh",
    "avail_discharge_energy_kwh",
    "aux_load_kw",
    "aux_load_pct",
    "cum_import_kwh",
    "cum_export_kwh",
    "revenue_cum_usd",
}


class EmsTelemetryImportError(FileFormatError):
    """Raised for file-level problems: bad format, missing columns, empty file."""


def _read_xlsx(content: bytes) -> tuple[list[str], list[list]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - surfaced to the caller as a clean import error
        raise EmsTelemetryImportError("Could not read the Excel file. Please upload a valid .xlsx workbook.",
                                      log_message=f"xlsx not readable: {type(exc).__name__}: {exc}") from exc

    sheet_name = next((n for n in wb.sheetnames if n.strip().lower() == "data"), None)
    if sheet_name is None:
        if len(wb.sheetnames) == 1:
            sheet_name = wb.sheetnames[0]
        else:
            raise EmsTelemetryImportError(
                f"Could not find a 'Data' worksheet. Available sheets: {', '.join(wb.sheetnames)}"
            )
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        raise EmsTelemetryImportError("Worksheet is empty")
    rows = [list(r) for r in rows_iter if any(v is not None for v in r)]
    return headers, rows


def _read_csv(content: bytes) -> tuple[list[str], list[list]]:
    try:
        text_content = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise EmsTelemetryImportError("Could not read the CSV file. Please save it as UTF-8 and upload it again.",
                                      log_message=f"csv not UTF-8: {exc}") from exc
    reader = csv.reader(io.StringIO(text_content))
    try:
        headers = [h.strip() for h in next(reader)]
    except StopIteration:
        raise EmsTelemetryImportError("File is empty")
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    return headers, rows


def _parse_timestamp(value) -> datetime:
    if isinstance(value, datetime):
        return value
    if value is None or str(value).strip() == "":
        raise ValueError("missing timestamp")
    return date_parser.parse(str(value).strip())


def _parse_decimal(value) -> Optional[Decimal]:
    if value is None:
        return None
    text_value = str(value).strip()
    if text_value == "":
        return None
    try:
        return Decimal(text_value)
    except InvalidOperation:
        raise ValueError(f"'{value}' is not a valid number")


def _parse_int(value) -> Optional[int]:
    if value is None:
        return None
    text_value = str(value).strip()
    if text_value == "":
        return None
    try:
        return int(float(text_value))
    except ValueError:
        raise ValueError(f"'{value}' is not a valid integer")


def _parse_text(value) -> Optional[str]:
    if value is None:
        return None
    text_value = str(value).strip()
    return text_value or None


def parse_ems_telemetry_file(
    filename: str, content: bytes, db: Session, ems_id: int
) -> tuple[list["schemas.EmsTelemetryCreate"], list["schemas.BulkEmsTelemetryRowError"], int]:
    """Parse an uploaded xlsx/csv file into validated ems_telemetry rows for a single EMS controller.

    The source file has no per-row controller identity column, so every row in the file
    is attributed to the given `ems_id`. Returns (valid_rows, row_errors, total_data_rows).
    Raises EmsTelemetryImportError for file-level problems.
    """
    lower_name = filename.lower()
    if lower_name.endswith((".xlsx", ".xls")):
        headers, raw_rows = _read_xlsx(content)
    elif lower_name.endswith(".csv"):
        headers, raw_rows = _read_csv(content)
    else:
        raise EmsTelemetryImportError("Unsupported file type. Please upload a .xlsx or .csv file.")

    if len(raw_rows) > MAX_ROWS:
        raise EmsTelemetryImportError(f"File has {len(raw_rows)} data rows, which exceeds the {MAX_ROWS} row limit.")

    normalized_headers = [h.strip().lower() for h in headers]
    header_index = {h: i for i, h in enumerate(normalized_headers) if h}

    ts_header = next((h for h in TIMESTAMP_HEADERS if h in header_index), None)
    if ts_header is None:
        raise EmsTelemetryImportError("Missing required column: ts (or Timestamp)")

    field_index = {field: header_index[header] for header, field in COLUMN_MAP.items() if header in header_index}
    ts_idx = header_index[ts_header]

    known_modes = {m.mode_code for m in db.query(models.EmsDispatchMode).all()}
    known_grid_services = {g.service_code for g in db.query(models.EmsGridService).all()}

    valid_rows: list[schemas.EmsTelemetryCreate] = []
    row_errors: list[schemas.BulkEmsTelemetryRowError] = []

    for i, raw in enumerate(raw_rows):
        row_number = i + 2  # header occupies row 1
        errors: list[str] = []

        def cell(idx: int):
            return raw[idx] if idx < len(raw) else None

        ts = None
        try:
            ts = _parse_timestamp(cell(ts_idx))
        except (ValueError, OverflowError) as exc:
            errors.append(f"Invalid timestamp: {exc}")

        field_values: dict = {}
        for field, idx in field_index.items():
            raw_value = cell(idx)
            try:
                if field in INT_FIELDS:
                    field_values[field] = _parse_int(raw_value)
                elif field in DECIMAL_FIELDS:
                    field_values[field] = _parse_decimal(raw_value)
                else:
                    field_values[field] = _parse_text(raw_value)
            except ValueError as exc:
                errors.append(f"{field}: {exc}")

        dispatch_mode = field_values.get("dispatch_mode")
        if dispatch_mode and dispatch_mode not in known_modes:
            errors.append(f"Unknown dispatch_mode '{dispatch_mode}'")
        grid_service = field_values.get("grid_service")
        if grid_service and grid_service not in known_grid_services:
            errors.append(f"Unknown grid_service '{grid_service}'")

        if errors:
            row_errors.append(schemas.BulkEmsTelemetryRowError(row=row_number, errors=errors))
            continue

        try:
            record = schemas.EmsTelemetryCreate(ems_id=ems_id, ts=ts, **field_values)
        except ValidationError as exc:
            messages = [f"{'.'.join(str(p) for p in e['loc'])}: {e['msg']}" for e in exc.errors()]
            row_errors.append(schemas.BulkEmsTelemetryRowError(row=row_number, errors=messages))
            continue

        valid_rows.append(record)

    return valid_rows, row_errors, len(raw_rows)


def upsert_ems_telemetry_rows(db: Session, rows: list["schemas.EmsTelemetryCreate"]) -> tuple[int, int]:
    """Insert new (ems_id, ts) rows and update rows that already exist.

    Existing rows are updated only in the columns the rows carry; the other columns keep their values.

    Returns (inserted_count, updated_count).
    """
    if not rows:
        return 0, 0

    dedup: dict[tuple, dict] = {}
    for r in rows:
        data = r.model_dump()
        dedup[(data["ems_id"], data["ts"])] = data
    all_rows = list(dedup.values())

    table = models.EmsTelemetry.__table__
    # A row that already exists (same ems_id + ts) is updated only in the columns the rows were given - for a
    # bulk file, the columns it has - so values a partial file doesn't carry keep what is stored.
    # (With no such column, ts = ts keeps it a no-op update, so the row still counts as updated.)
    sent = set().union(*(r.model_fields_set for r in rows))
    update_cols = [c.name for c in table.columns if c.name not in ("ems_id", "ts") and c.name in sent] or ["ts"]

    inserted = 0
    updated = 0
    chunk_size = 500
    for start in range(0, len(all_rows), chunk_size):
        chunk = all_rows[start : start + chunk_size]
        stmt = pg_insert(table).values(chunk)
        stmt = stmt.on_conflict_do_update(
            index_elements=["ems_id", "ts"],
            set_={name: getattr(stmt.excluded, name) for name in update_cols},
        )
        stmt = stmt.returning(literal_column("(xmax = 0)").label("was_insert"))
        result = db.execute(stmt)
        for row in result:
            if row.was_insert:
                inserted += 1
            else:
                updated += 1

    return inserted, updated
