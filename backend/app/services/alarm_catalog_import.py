"""Parse an alarm catalog file (.xlsx 'Alarm Catalog' sheet, or .csv) and upsert it by (source, alarm_code)."""
import csv
import io
from decimal import Decimal, InvalidOperation
from typing import Optional

import openpyxl
from pydantic import ValidationError
from sqlalchemy import func, literal_column
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

from bess_common.errors import FileFormatError
from app import models, schemas

MAX_ROWS = 10_000
SHEET_NAME = "alarm catalog"

# normalized header -> catalog field ('severity' may be a name or a level)
COLUMN_MAP = {
    "alarm_code": "alarm_code",
    "code": "alarm_code",
    "source": "source",
    "source_system": "source",
    "alarm_type": "alarm_type",
    "type": "alarm_type",
    "severity": "severity",
    "severity_level": "severity",
    "asset_type": "asset_type",
    "description": "description",
    "probable_cause": "probable_cause",
    "threshold": "threshold",
    "threshold_value": "threshold",
    "unit": "unit",
}
REQUIRED = ("alarm_code", "source", "alarm_type", "severity", "description")


class AlarmCatalogImportError(FileFormatError):
    """File-level problem: unreadable file, missing sheet or columns, too many rows."""


def _norm(header) -> str:
    return str(header or "").strip().lower().replace(" ", "_")


def _read_xlsx(content: bytes) -> tuple[str, list[str], list[list]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - reported to the caller as a clean import error
        raise AlarmCatalogImportError("Could not read the Excel file. Please upload a valid .xlsx workbook.",
                                      log_message=f"xlsx not readable: {type(exc).__name__}: {exc}") from exc
    name = next((n for n in wb.sheetnames if n.strip().lower() == SHEET_NAME), None)
    if name is None:
        if len(wb.sheetnames) != 1:
            raise AlarmCatalogImportError(
                f"Could not find an 'Alarm Catalog' worksheet. Available sheets: {', '.join(wb.sheetnames)}"
            )
        name = wb.sheetnames[0]
    rows_iter = wb[name].iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        raise AlarmCatalogImportError(f"Worksheet '{name}' is empty")
    rows = [list(r) for r in rows_iter if any(v is not None and str(v).strip() != "" for v in r)]
    return name, headers, rows


def _read_csv(content: bytes) -> tuple[None, list[str], list[list]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise AlarmCatalogImportError("Could not read the CSV file. Please save it as UTF-8 and upload it again.",
                                      log_message=f"csv not UTF-8: {exc}") from exc
    reader = csv.reader(io.StringIO(text))
    try:
        headers = [h.strip() for h in next(reader)]
    except StopIteration:
        raise AlarmCatalogImportError("File is empty")
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    return None, headers, rows


def _severity_level(value, by_name: dict[str, int]) -> int:
    """'Warning' / 'warning' / 1 / '1' -> 1."""
    if value is None or str(value).strip() == "":
        raise ValueError("severity is required")
    text = str(value).strip()
    if text.lower() in by_name:
        return by_name[text.lower()]
    try:
        level = int(float(text))
    except ValueError:
        raise ValueError(f"unknown severity '{text}' (use {', '.join(n.title() for n in by_name)} or 0-3)")
    if level not in by_name.values():
        raise ValueError(f"severity level {level} is not 0-3")
    return level


def _threshold(value) -> Optional[Decimal]:
    if value is None or str(value).strip() == "":
        return None
    try:
        return Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError(f"threshold '{value}' is not a number")


def parse_alarm_catalog_file(
    filename: str, content: bytes, db: Session
) -> tuple[Optional[str], list[schemas.AlarmCatalogCreate], list[schemas.BulkAlarmCatalogRowError], int, list[str]]:
    """Returns (sheet, valid rows, row errors, total rows, ignored columns)."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        sheet, headers, rows = _read_xlsx(content)
    elif lower.endswith(".csv"):
        sheet, headers, rows = _read_csv(content)
    else:
        raise AlarmCatalogImportError("Unsupported file type: upload an Excel (.xlsx) or CSV file")
    if len(rows) > MAX_ROWS:
        raise AlarmCatalogImportError(f"Too many rows ({len(rows)}); the limit is {MAX_ROWS}")

    columns: dict[int, str] = {}
    ignored: list[str] = []
    for i, header in enumerate(headers):
        field = COLUMN_MAP.get(_norm(header))
        if field and field not in columns.values():
            columns[i] = field
        elif header:
            ignored.append(header)
    missing = [f for f in REQUIRED if f not in columns.values()]
    if missing:
        raise AlarmCatalogImportError(
            f"Missing required column(s): {', '.join(missing)}. Found: {', '.join(h for h in headers if h)}"
        )

    by_name = {s.severity_name.lower(): s.severity_level for s in db.query(models.AlarmSeverity)}
    valid: list[schemas.AlarmCatalogCreate] = []
    errors: list[schemas.BulkAlarmCatalogRowError] = []
    first_row_of: dict[tuple[str, str], int] = {}

    for offset, raw in enumerate(rows):
        row_no = offset + 2  # 1-based, after the header row
        values = {field: (raw[i] if i < len(raw) else None) for i, field in columns.items()}
        problems: list[str] = []
        try:
            values["severity_level"] = _severity_level(values.pop("severity"), by_name)
        except ValueError as exc:
            problems.append(str(exc))
        try:
            values["threshold"] = _threshold(values.get("threshold"))
        except ValueError as exc:
            problems.append(str(exc))
        values = {k: (str(v) if isinstance(v, (int, float)) and k not in ("severity_level", "threshold") else v)
                  for k, v in values.items()}
        if not problems:
            try:
                item = schemas.AlarmCatalogCreate(**values)
            except ValidationError as exc:
                problems += [
                    f"{'.'.join(str(p) for p in e['loc']) or 'row'}: {e['msg'].replace('Value error, ', '')}"
                    for e in exc.errors()
                ]
            else:
                key = (item.source, item.alarm_code)
                if key in first_row_of:
                    problems.append(f"duplicate of row {first_row_of[key]} ({item.source} {item.alarm_code})")
                else:
                    first_row_of[key] = row_no
                    valid.append(item)
        if problems:
            errors.append(schemas.BulkAlarmCatalogRowError(row=row_no, errors=problems))
    return sheet, valid, errors, len(rows), ignored


def upsert_alarm_catalog(db: Session, rows: list[schemas.AlarmCatalogCreate]) -> tuple[int, int]:
    """Insert new (source, alarm_code) entries and update existing ones. Returns (inserted, updated)."""
    if not rows:
        return 0, 0
    table = models.AlarmCatalog.__table__
    data = [r.model_dump() for r in rows]
    stmt = pg_insert(table).values(data)
    update_cols = [c for c in data[0] if c not in ("source", "alarm_code")]
    stmt = stmt.on_conflict_do_update(
        constraint="alarm_catalog_source_code_key",
        set_={**{c: getattr(stmt.excluded, c) for c in update_cols}, "updated_at": func.now()},
    ).returning(literal_column("(xmax = 0)").label("was_insert"))
    inserted = updated = 0
    for row in db.execute(stmt):
        if row.was_insert:
            inserted += 1
        else:
            updated += 1
    return inserted, updated
