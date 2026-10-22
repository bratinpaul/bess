"""Bulk load of site demand readings (demand_data.csv / .xlsx) into battery_telemetry.demand_telemetry.

Columns: ts (or timestamp), metering_point_code (or mp_code; optional when one metering point is chosen for the whole
file), gross_load_kw, behind_meter_gen_kw, net_load_kw, load_reactive_kvar. Header case and surrounding spaces are
ignored; unknown columns are reported and skipped.

Row validation: a parseable timestamp, a known metering point, numbers in range (gross and generation >= 0),
net = gross - generation within 1 kW (derived when the net column is empty), and no duplicate (point, ts) in the file
(the last one wins). Existing (metering point, ts) readings are updated, new ones inserted. Timestamps without a zone
are plant time (Asia/Kolkata).
"""
import csv
import io
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Optional
from zoneinfo import ZoneInfo

import openpyxl
from dateutil import parser as date_parser
from pydantic import ValidationError
from sqlalchemy import literal_column
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

from app import models, schemas

MAX_ROWS = 300_000
PLANT_TZ = ZoneInfo("Asia/Kolkata")
TIMESTAMP_HEADERS = ("ts", "timestamp", "time")
CODE_HEADERS = ("metering_point_code", "mp_code", "metering_point")
VALUE_FIELDS = ("gross_load_kw", "behind_meter_gen_kw", "net_load_kw", "load_reactive_kvar")


class DemandTelemetryImportError(Exception):
    """File-level problems: bad format, missing columns, empty file."""


def _read_xlsx(content: bytes) -> tuple[list[str], list[list]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - surfaced as a clean import error
        raise DemandTelemetryImportError(f"Could not read Excel file: {exc}") from exc
    name = next((n for n in wb.sheetnames if n.strip().lower() in ("data", "demand", "demand_data")), wb.sheetnames[0])
    rows_iter = wb[name].iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        raise DemandTelemetryImportError("Worksheet is empty")
    return headers, [list(r) for r in rows_iter if any(v is not None for v in r)]


def _read_csv(content: bytes) -> tuple[list[str], list[list]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DemandTelemetryImportError(f"Could not read CSV file as UTF-8: {exc}") from exc
    reader = csv.reader(io.StringIO(text))
    try:
        headers = [h.strip() for h in next(reader)]
    except StopIteration:
        raise DemandTelemetryImportError("File is empty")
    return headers, [row for row in reader if any(cell.strip() for cell in row)]


def _ts(value) -> datetime:
    if isinstance(value, datetime):
        dt = value
    elif value is None or str(value).strip() == "":
        raise ValueError("missing timestamp")
    else:
        dt = date_parser.parse(str(value).strip())
    return dt if dt.tzinfo else dt.replace(tzinfo=PLANT_TZ)


def _dec(value) -> Optional[Decimal]:
    if value is None or str(value).strip() == "":
        return None
    try:
        d = Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError(f"'{value}' is not a number")
    if not d.is_finite():
        raise ValueError(f"'{value}' is not a finite number")
    return d


def parse_demand_file(filename: str, content: bytes, db: Session, default_mp_id: Optional[int] = None):
    """Returns (valid rows, row errors, total data rows, unknown columns, metering point codes seen)."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        headers, raw_rows = _read_xlsx(content)
    elif name.endswith(".csv"):
        headers, raw_rows = _read_csv(content)
    else:
        raise DemandTelemetryImportError("Unsupported file type. Please upload a .csv or .xlsx file.")
    if len(raw_rows) > MAX_ROWS:
        raise DemandTelemetryImportError(f"File has {len(raw_rows)} rows, above the {MAX_ROWS} row limit.")

    norm = [h.strip().lower() for h in headers]
    idx = {h: i for i, h in enumerate(norm) if h}
    ts_h = next((h for h in TIMESTAMP_HEADERS if h in idx), None)
    if ts_h is None:
        raise DemandTelemetryImportError("Missing required column: ts (or timestamp)")
    code_h = next((h for h in CODE_HEADERS if h in idx), None)
    if code_h is None and default_mp_id is None:
        raise DemandTelemetryImportError("The file has no metering_point_code column: choose the metering point for it")
    present = [f for f in VALUE_FIELDS if f in idx]
    if not ({"gross_load_kw", "net_load_kw"} & set(present)):
        raise DemandTelemetryImportError("Missing columns: need gross_load_kw or net_load_kw")
    known = {ts_h, code_h, *present}
    ignored = [h for h, n in zip(headers, norm) if n and n not in known]

    points = {mp.mp_code.upper(): mp.metering_point_id for mp in db.query(models.MeteringPoint)}
    valid, errors, seen_codes = {}, [], set()
    for i, raw in enumerate(raw_rows):
        row_no = i + 2
        problems: list[str] = []
        cell = lambda j: raw[j] if j < len(raw) else None  # noqa: E731
        try:
            ts = _ts(cell(idx[ts_h]))
        except (ValueError, OverflowError) as exc:
            problems.append(f"Invalid timestamp: {exc}")
            ts = None
        mp_id = default_mp_id
        if code_h is not None:
            code = str(cell(idx[code_h]) or "").strip()
            if code:
                mp_id = points.get(code.upper())
                seen_codes.add(code)
                if mp_id is None:
                    problems.append(f"Unknown metering point '{code}'")
            elif default_mp_id is None:
                problems.append("Missing metering_point_code")
        values = {}
        for f in present:
            try:
                values[f] = _dec(cell(idx[f]))
            except ValueError as exc:
                problems.append(f"{f}: {exc}")
        if problems:
            errors.append(schemas.BulkDemandTelemetryRowError(row=row_no, errors=problems))
            continue
        try:
            rec = schemas.DemandTelemetryCreate(metering_point_id=mp_id, ts=ts, **values)
        except ValidationError as exc:
            msgs = [(".".join(str(p) for p in e["loc"]) + ": " if e["loc"] else "") + e["msg"].replace("Value error, ", "")
                    for e in exc.errors()]
            errors.append(schemas.BulkDemandTelemetryRowError(row=row_no, errors=msgs))
            continue
        valid[(mp_id, ts)] = rec  # a later duplicate in the file wins
    return list(valid.values()), errors, len(raw_rows), ignored, sorted(seen_codes)


def upsert_demand_rows(db: Session, rows: list["schemas.DemandTelemetryCreate"]) -> tuple[int, int]:
    """Insert new (metering point, ts) readings and update existing ones. Returns (inserted, updated)."""
    if not rows:
        return 0, 0
    table = models.DemandTelemetry.__table__
    data = [r.model_dump() for r in rows]
    inserted = updated = 0
    for start in range(0, len(data), 1000):
        stmt = pg_insert(table).values(data[start:start + 1000])
        stmt = stmt.on_conflict_do_update(
            index_elements=["metering_point_id", "ts"],
            set_={f: getattr(stmt.excluded, f) for f in VALUE_FIELDS},
        ).returning(literal_column("(xmax = 0)").label("was_insert"))
        for r in db.execute(stmt):
            if r.was_insert:
                inserted += 1
            else:
                updated += 1
    return inserted, updated
