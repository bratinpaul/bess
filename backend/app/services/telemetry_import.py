import csv
import io
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Optional

import openpyxl
from dateutil import parser as date_parser
from pydantic import ValidationError
from sqlalchemy import literal_column
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

from bess_common.errors import FileFormatError
from app import models, schemas

MAX_ROWS = 200_000

REQUIRED_HEADERS = {"system_name", "timestamp"}

# Maps a normalized (lowercased) source column header to a system_telemetry field.
COLUMN_MAP = {
    "mode": "mode_code",
    # SOC as reported by the BMS: the original SOC_% header and the new column name
    "soc_%": "soc_bms_pct",
    "soc_bms_%": "soc_bms_pct",
    "soc_bms_pct": "soc_bms_pct",
    "soh_%": "soh_pct",
    "pack_voltage_v": "pack_voltage_v",
    "pack_current_a": "pack_current_a",
    "power_kw": "power_kw",
    "c_rate": "c_rate",
    "cell_v_avg_v": "cell_v_avg_v",
    "cell_v_min_v": "cell_v_min_v",
    "cell_v_max_v": "cell_v_max_v",
    "cell_v_delta_mv": "cell_v_delta_mv",
    "cell_temp_avg_c": "cell_temp_avg_c",
    "cell_temp_min_c": "cell_temp_min_c",
    "cell_temp_max_c": "cell_temp_max_c",
    "coolant_temp_c": "coolant_temp_c",
    "ambient_temp_c": "ambient_temp_c",
    "insulation_res_kohm": "insulation_res_kohm",
    "avail_charge_kw": "avail_charge_kw",
    "avail_discharge_kw": "avail_discharge_kw",
    "cum_charge_kwh": "cum_charge_kwh",
    "cum_discharge_kwh": "cum_discharge_kwh",
    "equiv_full_cycles": "equiv_full_cycles",
    "contactor_status": "contactor_state",
    "alarm_level": "alarm_level",
    # current sensing and SOC estimation (headers as the column names; "_%" spelling accepted for the SOCs)
    "measured_current_a": "measured_current_a",
    "current_sensor_offset_a": "current_sensor_offset_a",
    "soc_coulomb_%": "soc_coulomb_pct",
    "soc_coulomb_pct": "soc_coulomb_pct",
    "soc_battery_telemetry_%": "soc_battery_telemetry_pct",
    "soc_battery_telemetry_pct": "soc_battery_telemetry_pct",
    "usable_capacity_ah": "usable_capacity_ah",
    "coulombic_efficiency": "coulombic_efficiency",
    "at_rest": "at_rest",
    "rest_duration_min": "rest_duration_min",
    "ocv_rested_v": "ocv_rested_v",
}

BOOL_FIELDS = {"at_rest"}

DECIMAL_FIELDS = {
    "soc_bms_pct",
    "soh_pct",
    "pack_voltage_v",
    "pack_current_a",
    "power_kw",
    "c_rate",
    "cell_v_avg_v",
    "cell_v_min_v",
    "cell_v_max_v",
    "cell_v_delta_mv",
    "cell_temp_avg_c",
    "cell_temp_min_c",
    "cell_temp_max_c",
    "coolant_temp_c",
    "ambient_temp_c",
    "insulation_res_kohm",
    "avail_charge_kw",
    "avail_discharge_kw",
    "cum_charge_kwh",
    "cum_discharge_kwh",
    "equiv_full_cycles",
    "measured_current_a",
    "current_sensor_offset_a",
    "soc_coulomb_pct",
    "soc_battery_telemetry_pct",
    "usable_capacity_ah",
    "coulombic_efficiency",
    "rest_duration_min",
    "ocv_rested_v",
}


class TelemetryImportError(FileFormatError):
    """Raised for file-level problems: bad format, missing columns, empty file."""


def _read_xlsx(content: bytes) -> tuple[list[str], list[list]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - surfaced to the caller as a clean import error
        raise TelemetryImportError("Could not read the Excel file. Please upload a valid .xlsx workbook.",
                                   log_message=f"xlsx not readable: {type(exc).__name__}: {exc}") from exc

    sheet_name = next((n for n in wb.sheetnames if n.strip().lower() == "data"), None)
    if sheet_name is None:
        if len(wb.sheetnames) == 1:
            sheet_name = wb.sheetnames[0]
        else:
            raise TelemetryImportError(
                f"Could not find a 'Data' worksheet. Available sheets: {', '.join(wb.sheetnames)}"
            )
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        raise TelemetryImportError("Worksheet is empty")
    rows = [list(r) for r in rows_iter if any(v is not None for v in r)]
    return headers, rows


def _read_csv(content: bytes) -> tuple[list[str], list[list]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise TelemetryImportError("Could not read the CSV file. Please save it as UTF-8 and upload it again.",
                                   log_message=f"csv not UTF-8: {exc}") from exc
    reader = csv.reader(io.StringIO(text))
    try:
        headers = [h.strip() for h in next(reader)]
    except StopIteration:
        raise TelemetryImportError("File is empty")
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    return headers, rows


def _parse_timestamp(value) -> datetime:
    if isinstance(value, datetime):
        return value
    if value is None or str(value).strip() == "":
        raise ValueError("missing timestamp")
    return date_parser.parse(str(value).strip())


def _parse_decimal(value) -> Optional[Decimal]:
    if value is None:
        return None
    text = str(value).strip()
    if text == "":
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        raise ValueError(f"'{value}' is not a valid number")


def _parse_int(value) -> Optional[int]:
    if value is None:
        return None
    text = str(value).strip()
    if text == "":
        return None
    try:
        return int(float(text))
    except ValueError:
        raise ValueError(f"'{value}' is not a valid integer")


def _parse_bool(value) -> Optional[bool]:
    if value is None or isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text == "":
        return None
    if text in ("true", "t", "yes", "y", "1", "1.0"):
        return True
    if text in ("false", "f", "no", "n", "0", "0.0"):
        return False
    raise ValueError(f"'{value}' is not true / false")


def _parse_text(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def parse_telemetry_file(
    filename: str, content: bytes, db: Session
) -> tuple[list["schemas.SystemTelemetryCreate"], list["schemas.BulkTelemetryRowError"], int]:
    """Parse an uploaded xlsx/csv file into validated rows.

    Returns (valid_rows, row_errors, total_data_rows). Raises TelemetryImportError
    for file-level problems (wrong type, missing required columns, unreadable file).
    """
    lower_name = filename.lower()
    if lower_name.endswith((".xlsx", ".xls")):
        headers, raw_rows = _read_xlsx(content)
    elif lower_name.endswith(".csv"):
        headers, raw_rows = _read_csv(content)
    else:
        raise TelemetryImportError("Unsupported file type. Please upload a .xlsx or .csv file.")

    if len(raw_rows) > MAX_ROWS:
        raise TelemetryImportError(f"File has {len(raw_rows)} data rows, which exceeds the {MAX_ROWS} row limit.")

    normalized_headers = [h.strip().lower() for h in headers]
    header_index = {h: i for i, h in enumerate(normalized_headers) if h}

    missing_required = REQUIRED_HEADERS - set(header_index)
    if missing_required:
        raise TelemetryImportError(f"Missing required column(s): {', '.join(sorted(missing_required))}")

    field_index = {field: header_index[header] for header, field in COLUMN_MAP.items() if header in header_index}
    system_name_idx = header_index["system_name"]
    timestamp_idx = header_index["timestamp"]

    systems = {s.system_name.strip().lower(): s.battery_system_id for s in db.query(models.BatterySystem).all()}
    known_modes = {m.mode_code for m in db.query(models.OperatingMode).all()}
    known_contactor_states = {c.state_code for c in db.query(models.ContactorState).all()}

    valid_rows: list[schemas.SystemTelemetryCreate] = []
    row_errors: list[schemas.BulkTelemetryRowError] = []

    for i, raw in enumerate(raw_rows):
        row_number = i + 2  # header occupies row 1
        errors: list[str] = []

        def cell(idx: int):
            return raw[idx] if idx < len(raw) else None

        system_name = _parse_text(cell(system_name_idx))
        system_id = None
        if not system_name:
            errors.append("System_Name is required")
        else:
            system_id = systems.get(system_name.lower())
            if system_id is None:
                errors.append(f"Unknown System_Name '{system_name}' (no matching battery system)")

        ts = None
        try:
            ts = _parse_timestamp(cell(timestamp_idx))
        except (ValueError, OverflowError) as exc:
            errors.append(f"Invalid Timestamp: {exc}")

        field_values: dict = {}
        for field, idx in field_index.items():
            raw_value = cell(idx)
            try:
                if field == "alarm_level":
                    field_values[field] = _parse_int(raw_value)
                elif field in BOOL_FIELDS:
                    field_values[field] = _parse_bool(raw_value)
                elif field in DECIMAL_FIELDS:
                    field_values[field] = _parse_decimal(raw_value)
                else:
                    field_values[field] = _parse_text(raw_value)
            except ValueError as exc:
                errors.append(f"{field}: {exc}")

        mode_code = field_values.get("mode_code")
        if mode_code and mode_code not in known_modes:
            errors.append(f"Unknown Mode '{mode_code}'")
        contactor_state = field_values.get("contactor_state")
        if contactor_state and contactor_state not in known_contactor_states:
            errors.append(f"Unknown Contactor_Status '{contactor_state}'")

        if errors:
            row_errors.append(schemas.BulkTelemetryRowError(row=row_number, errors=errors))
            continue

        try:
            record = schemas.SystemTelemetryCreate(system_id=system_id, ts=ts, **field_values)
        except ValidationError as exc:
            messages = [f"{'.'.join(str(p) for p in e['loc'])}: {e['msg']}" for e in exc.errors()]
            row_errors.append(schemas.BulkTelemetryRowError(row=row_number, errors=messages))
            continue

        valid_rows.append(record)

    return valid_rows, row_errors, len(raw_rows)


def upsert_telemetry_rows(db: Session, rows: list["schemas.SystemTelemetryCreate"]) -> tuple[int, int]:
    """Insert new (system_id, ts) rows and update rows that already exist.

    Existing rows are updated only in the columns the rows carry; the other columns keep their values.

    Returns (inserted_count, updated_count). Uses Postgres's `xmax = 0` trick on the
    RETURNING clause to tell inserts from updates, since comparing naive vs.
    timezone-aware datetimes ourselves would be unreliable.
    """
    if not rows:
        return 0, 0

    # De-duplicate within the file itself: keep the last occurrence of a given key.
    dedup: dict[tuple, dict] = {}
    for r in rows:
        data = r.model_dump()
        dedup[(data["system_id"], data["ts"])] = data
    all_rows = list(dedup.values())

    table = models.SystemTelemetry.__table__
    # A row that already exists (same system_id + ts) is updated only in the columns the rows were given - for a
    # bulk file, the columns it has - so values a partial file doesn't carry keep what is stored.
    # (With no such column, ts = ts keeps it a no-op update, so the row still counts as updated.)
    sent = set().union(*(r.model_fields_set for r in rows))
    update_cols = [c.name for c in table.columns if c.name not in ("system_id", "ts") and c.name in sent] or ["ts"]

    inserted = 0
    updated = 0
    chunk_size = 500
    for start in range(0, len(all_rows), chunk_size):
        chunk = all_rows[start : start + chunk_size]
        stmt = pg_insert(table).values(chunk)
        stmt = stmt.on_conflict_do_update(
            index_elements=["system_id", "ts"],
            set_={name: getattr(stmt.excluded, name) for name in update_cols},
        )
        stmt = stmt.returning(literal_column("(xmax = 0)").label("was_insert"))
        result = db.execute(stmt)
        for row in result:
            if row.was_insert:
                inserted += 1
            else:
                updated += 1

    return inserted, updated
