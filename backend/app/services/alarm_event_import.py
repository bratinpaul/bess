"""Parse an alarm log (.xlsx 'Alarms' sheet, or .csv) and upsert it into alarm_event by (alarm_id, start time).

Per row: the (source, alarm_code) must be in the alarm catalog; the reported asset code is matched to its
master-data record (services/asset_matcher.py) and stored under that record's type and name; times and
the acknowledgement must be consistent.
Columns that the catalog or the timestamps already give (alarm_type, description, probable_cause,
duration_min, status, start_date) aren't stored; status and duration_min are checked when present.

An alarm that already exists is updated only in the columns the file has, so a partial file (e.g. a plant
snapshot without the acknowledgement) doesn't blank the values already stored.
"""
import csv
import io
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Optional

import openpyxl
from dateutil import parser as date_parser
from sqlalchemy import func, literal_column
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

from bess_common.errors import FileFormatError
from app import models, schemas
from app.services.asset_matcher import LABELS, AssetMatcher, AssetMatchError

MAX_ROWS = 100_000
SHEET_NAMES = ("alarms", "alarm log", "alarm events", "data")

COLUMN_MAP = {
    "alarm_id": "alarm_id", "id": "alarm_id",
    "source_system": "source", "source": "source",
    "asset_type": "asset_type",
    "asset_code": "asset_code", "asset": "asset_code",
    "alarm_code": "alarm_code", "code": "alarm_code",
    "severity": "severity",
    "severity_level": "severity_level",
    "start_time": "start_time", "time": "start_time", "start": "start_time",
    "end_time": "end_time", "end": "end_time",
    "trigger_value": "trigger_value",
    "threshold_value": "threshold_value", "threshold": "threshold_value",
    "unit": "unit",
    "acknowledged_by": "acknowledged_by",
    "acknowledged_at": "acknowledged_at",
    # checked, not stored
    "status": "status",
    "duration_min": "duration_min",
}
# given by the catalog or derived from the times: accepted and skipped
KNOWN_UNSTORED = {"alarm_type", "description", "probable_cause", "start_date"}
REQUIRED = ("alarm_id", "source", "alarm_code", "start_time")

# file field -> the table columns it sets; a column is updated on existing rows only when its field is in the file
FIELD_COLUMNS = {
    "alarm_code": ("alarm_catalog_id",),
    "severity": ("severity_level",),
    "severity_level": ("severity_level",),
    "asset_code": ("asset_type", "asset_code", "linked_asset_type", "linked_asset_id", "asset_match", "battery_system_id"),
    "end_time": ("end_time",),
    "trigger_value": ("trigger_value",),
    "threshold_value": ("threshold_value",),
    "unit": ("unit",),
    "acknowledged_by": ("acknowledged_by",),
    "acknowledged_at": ("acknowledged_at",),
}


class AlarmEventImportError(FileFormatError):
    """File-level problem: unreadable file, missing sheet or columns, too many rows."""


def _norm(header) -> str:
    return str(header or "").strip().lower().replace(" ", "_")


def _read_xlsx(content: bytes) -> tuple[str, list[str], list[list]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - reported to the caller as a clean import error
        raise AlarmEventImportError("Could not read the Excel file. Please upload a valid .xlsx workbook.",
                                    log_message=f"xlsx not readable: {type(exc).__name__}: {exc}") from exc
    lowered = {n.strip().lower(): n for n in wb.sheetnames}
    name = next((lowered[s] for s in SHEET_NAMES if s in lowered), None)
    if name is None:
        if len(wb.sheetnames) != 1:
            raise AlarmEventImportError(f"Could not find an 'Alarms' worksheet. Available sheets: {', '.join(wb.sheetnames)}")
        name = wb.sheetnames[0]
    rows_iter = wb[name].iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        raise AlarmEventImportError(f"Worksheet '{name}' is empty")
    rows = [list(r) for r in rows_iter if any(v is not None and str(v).strip() != "" for v in r)]
    return name, headers, rows


def _read_csv(content: bytes) -> tuple[None, list[str], list[list]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise AlarmEventImportError("Could not read the CSV file. Please save it as UTF-8 and upload it again.",
                                    log_message=f"csv not UTF-8: {exc}") from exc
    reader = csv.reader(io.StringIO(text))
    try:
        headers = [h.strip() for h in next(reader)]
    except StopIteration:
        raise AlarmEventImportError("File is empty")
    return None, headers, [row for row in reader if any(cell.strip() for cell in row)]


def _text(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _time(value, label: str) -> Optional[datetime]:
    """No time-zone conversion: naive times are taken as plant local time, like the telemetry uploads."""
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value
    try:
        return date_parser.parse(str(value).strip())
    except (ValueError, OverflowError):
        raise ValueError(f"{label} '{value}' is not a date/time")


def _number(value, label: str) -> Optional[Decimal]:
    if value is None or str(value).strip() == "":
        return None
    try:
        return Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError(f"{label} '{value}' is not a number")


def parse_alarm_event_file(filename: str, content: bytes, db: Session):
    """Returns (sheet, rows ready for upsert, row errors, total rows, ignored columns, assets linked, columns to update)."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        sheet, headers, rows = _read_xlsx(content)
    elif lower.endswith(".csv"):
        sheet, headers, rows = _read_csv(content)
    else:
        raise AlarmEventImportError("Unsupported file type: upload an Excel (.xlsx) or CSV file")
    if len(rows) > MAX_ROWS:
        raise AlarmEventImportError(f"Too many rows ({len(rows)}); the limit is {MAX_ROWS}")

    columns: dict[int, str] = {}
    ignored: list[str] = []
    for i, header in enumerate(headers):
        key = _norm(header)
        field = COLUMN_MAP.get(key)
        if field and field not in columns.values():
            columns[i] = field
        elif header and key not in KNOWN_UNSTORED:
            ignored.append(header)
    missing = [f for f in REQUIRED if f not in columns.values()]
    if missing:
        raise AlarmEventImportError(
            f"Missing required column(s): {', '.join(missing)}. Found: {', '.join(h for h in headers if h)}"
        )

    present = set(columns.values())
    update_cols = sorted({col for field in present for col in FIELD_COLUMNS.get(field, ())})

    catalog = {(c.source, c.alarm_code): c for c in db.query(models.AlarmCatalog)}
    by_severity = {s.severity_name.lower(): s.severity_level for s in db.query(models.AlarmSeverity)}
    matcher = AssetMatcher(db)
    existing = {a: t for a, t in db.query(models.AlarmEvent.alarm_id, models.AlarmEvent.time)}

    valid: list[dict] = []
    errors: list[schemas.BulkAlarmEventRowError] = []
    seen: dict[str, int] = {}
    linked = 0

    for offset, raw in enumerate(rows):
        row_no = offset + 2  # 1-based, after the header row
        v = {field: (raw[i] if i < len(raw) else None) for i, field in columns.items()}
        problems: list[str] = []
        alarm_id = _text(v.get("alarm_id"))

        def attempt(fn, *args):
            try:
                return fn(*args)
            except ValueError as exc:
                problems.append(str(exc))
                return None

        start = attempt(_time, v.get("start_time"), "start_time")
        end = attempt(_time, v.get("end_time"), "end_time")
        ack_at = attempt(_time, v.get("acknowledged_at"), "acknowledged_at")
        trigger = attempt(_number, v.get("trigger_value"), "trigger_value")
        threshold = attempt(_number, v.get("threshold_value"), "threshold_value")
        duration = attempt(_number, v.get("duration_min"), "duration_min")
        ack_by = _text(v.get("acknowledged_by"))
        source = (_text(v.get("source")) or "").upper()
        code = _text(v.get("alarm_code"))

        if not alarm_id:
            problems.append("alarm_id is required")
        elif alarm_id in seen:
            problems.append(f"duplicate of row {seen[alarm_id]} (alarm_id {alarm_id})")
        if _text(v.get("start_time")) is None:  # a start time that doesn't parse is reported above
            problems.append("start_time is required")

        entry = catalog.get((source, code))
        if not code:
            problems.append("alarm_code is required")
        elif entry is None:
            problems.append(f"{source or '?'} alarm '{code}' is not in the alarm catalog")

        # severity: level column, else name, else the catalog's
        level = None
        level_text, name_text = _text(v.get("severity_level")), _text(v.get("severity"))
        if level_text is not None:
            try:
                level = int(float(level_text))
                if level not in by_severity.values():
                    problems.append(f"severity_level {level} is not 0-3")
            except ValueError:
                problems.append(f"severity_level '{level_text}' is not a number")
        if name_text is not None:
            named = by_severity.get(name_text.lower())
            if named is None:
                problems.append(f"unknown severity '{name_text}'")
            elif level is not None and named != level:
                problems.append(f"severity '{name_text}' doesn't match severity_level {level}")
            elif level is None:
                level = named
        if level is None and entry is not None:
            level = entry.severity_level

        if start and end and end < start:
            problems.append("end_time is before start_time")
        if (ack_by is None) != (ack_at is None):
            problems.append("acknowledged_by and acknowledged_at must be given together")
        if start and ack_at and ack_at < start:
            problems.append("acknowledged_at is before start_time")
        status = (_text(v.get("status")) or "").lower()
        if status == "active" and end is not None:
            problems.append("status is Active but end_time is set")
        if status == "cleared" and end is None:
            problems.append("status is Cleared but end_time is empty")
        if status and status not in ("active", "cleared"):
            problems.append(f"unknown status '{v.get('status')}' (Active or Cleared)")
        if duration is not None and start and end and abs(float(duration) - (end - start).total_seconds() / 60) > 1:
            problems.append(f"duration_min {duration} doesn't match start/end ({(end - start).total_seconds() / 60:.0f} min)")

        if alarm_id and start and alarm_id in existing and existing[alarm_id].replace(tzinfo=None) != start.replace(tzinfo=None):
            problems.append(f"alarm_id {alarm_id} already exists with start time {existing[alarm_id]:%Y-%m-%d %H:%M}")

        asset_type, asset_code = _text(v.get("asset_type")), _text(v.get("asset_code"))
        match = None
        if asset_code:
            try:
                match = matcher.match(asset_code, asset_type, source)
            except AssetMatchError as exc:
                problems.append(str(exc))

        if problems:
            errors.append(schemas.BulkAlarmEventRowError(row=row_no, alarm_id=alarm_id, errors=problems))
            continue
        seen[alarm_id] = row_no
        linked += match is not None
        valid.append(
            dict(
                alarm_id=alarm_id,
                time=start,
                alarm_catalog_id=entry.alarm_catalog_id,
                severity_level=level,
                # a matched asset is stored under its master-data type and name (PACK-01 -> 01-01-01-PACK-01)
                asset_type=LABELS[match.asset_type] if match else asset_type,
                asset_code=match.name if match else asset_code,
                linked_asset_type=match.asset_type if match else None,
                linked_asset_id=match.asset_id if match else None,
                asset_match=match.method if match else None,
                battery_system_id=match.battery_system_id if match else None,
                end_time=end,
                trigger_value=trigger,
                threshold_value=threshold,
                unit=_text(v.get("unit")),
                acknowledged_by=ack_by,
                acknowledged_at=ack_at,
            )
        )
    return sheet, valid, errors, len(rows), ignored, linked, update_cols


def upsert_alarm_events(db: Session, rows: list[dict], update_cols: Optional[list[str]] = None) -> tuple[int, int]:
    """Insert new (alarm_id, time) rows and update existing ones. Returns (inserted, updated).

    `update_cols` limits what an existing row gets updated with (the columns the file had); None updates all.
    """
    if not rows:
        return 0, 0
    table = models.AlarmEvent.__table__
    inserted = updated = 0
    for start in range(0, len(rows), 500):
        chunk = rows[start : start + 500]
        stmt = pg_insert(table).values(chunk)
        cols = [c for c in chunk[0] if c not in ("alarm_id", "time") and (update_cols is None or c in update_cols)]
        stmt = stmt.on_conflict_do_update(
            index_elements=["alarm_id", "time"],
            set_={**{c: getattr(stmt.excluded, c) for c in cols}, "updated_at": func.now()},
        ).returning(literal_column("(xmax = 0)").label("was_insert"))
        for row in db.execute(stmt):
            if row.was_insert:
                inserted += 1
            else:
                updated += 1
    return inserted, updated
