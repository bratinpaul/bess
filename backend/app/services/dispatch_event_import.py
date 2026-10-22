"""Parse an EMS dispatch log (.xlsx 'Dispatch Events' sheet, or .csv) and upsert it into dispatch_event
by (dispatch_id, effective_from).

Per row: the EMS, dispatch mode and grid service must exist in master data; times, SOC and energy must be
consistent; and one EMS runs one dispatch at a time, so windows on the same EMS may not overlap (checked
against the rest of the file and the dispatches already stored). The exception is a running dispatch (no end
yet): a newer dispatch on the same EMS ends it, as a new command does on the EMS - it is closed (Completed)
at the new dispatch's start.
Columns the other columns already give (duration_min, net_energy_kwh, issued_date) aren't stored;
duration_min and net_energy_kwh are checked when present.

A dispatch that already exists is updated only in the columns the file has, so a partial file (e.g. a
plant snapshot without energy or revenue) doesn't blank the values already stored.
"""
import csv
import io
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Optional

import openpyxl
from dateutil import parser as date_parser
from sqlalchemy import func, literal_column
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

from app import models, schemas
from app.schemas.dispatch_event import COMMAND_CODES, DISPATCH_STATUSES

MAX_ROWS = 100_000
SHEET_NAMES = ("dispatch events", "dispatch_events", "dispatch event log", "dispatch", "data")
# rounding in the source: net energy may differ from import - export by a fraction of a kWh
NET_TOLERANCE_KWH = Decimal("0.5")

COLUMN_MAP = {
    "dispatch_id": "dispatch_id", "id": "dispatch_id",
    "ems_code": "ems_code", "ems": "ems_code",
    "command_code": "command_code", "command": "command_code",
    "dispatch_mode": "dispatch_mode", "mode": "dispatch_mode",
    "grid_service": "grid_service", "service": "grid_service",
    "source": "source", "issued_by": "source",
    "issued_at": "issued_at",
    "effective_from": "effective_from", "start_time": "effective_from", "time": "effective_from",
    "effective_to": "effective_to", "end_time": "effective_to",
    "p_setpoint_kw": "p_setpoint_kw", "setpoint_kw": "p_setpoint_kw",
    "target_soc_pct": "target_soc_pct",
    "soc_start_pct": "soc_start_pct",
    "soc_end_pct": "soc_end_pct",
    "ramp_rate_kw_s": "ramp_rate_kw_s",
    "import_kwh": "import_kwh",
    "export_kwh": "export_kwh",
    "avg_price_mwh": "avg_price_mwh",
    "revenue_usd": "revenue_usd",
    "accepted": "accepted",
    "status": "status",
    "note": "note",
    # checked, not stored
    "duration_min": "duration_min",
    "net_energy_kwh": "net_energy_kwh",
}
# derived from issued_at: accepted and skipped
KNOWN_UNSTORED = {"issued_date"}
REQUIRED = ("dispatch_id", "ems_code", "command_code", "dispatch_mode", "effective_from")

# file field -> the table columns it sets; a column is updated on existing rows only when its field is in the file
FIELD_COLUMNS = {
    "ems_code": ("ems_id", "battery_system_id"),
    "command_code": ("command_code",),
    "dispatch_mode": ("dispatch_mode",),
    "grid_service": ("grid_service",),
    "source": ("source",),
    "issued_at": ("issued_at",),
    # the end time also decides the status when there is no status column
    "effective_to": ("effective_to", "status"),
    "p_setpoint_kw": ("p_setpoint_kw",),
    "target_soc_pct": ("target_soc_pct",),
    "soc_start_pct": ("soc_start_pct",),
    "soc_end_pct": ("soc_end_pct",),
    "ramp_rate_kw_s": ("ramp_rate_kw_s",),
    "import_kwh": ("import_kwh",),
    "export_kwh": ("export_kwh",),
    "avg_price_mwh": ("avg_price_mwh",),
    "revenue_usd": ("revenue_usd",),
    "accepted": ("accepted",),
    "status": ("status",),
    "note": ("note",),
}

TRUE_WORDS = {"true", "yes", "y", "1", "=true()"}
FALSE_WORDS = {"false", "no", "n", "0", "=false()"}


class DispatchEventImportError(Exception):
    """File-level problem: unreadable file, missing sheet or columns, too many rows."""


def _norm(header) -> str:
    return str(header or "").strip().lower().replace(" ", "_")


def _read_xlsx(content: bytes) -> tuple[str, list[str], list[list]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - reported to the caller as a clean import error
        raise DispatchEventImportError(f"Could not read Excel file: {exc}") from exc
    lowered = {n.strip().lower(): n for n in wb.sheetnames}
    name = next((lowered[s] for s in SHEET_NAMES if s in lowered), None)
    if name is None:
        if len(wb.sheetnames) != 1:
            raise DispatchEventImportError(
                f"Could not find a 'Dispatch Events' worksheet. Available sheets: {', '.join(wb.sheetnames)}")
        name = wb.sheetnames[0]
    rows_iter = wb[name].iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        raise DispatchEventImportError(f"Worksheet '{name}' is empty")
    rows = [list(r) for r in rows_iter if any(v is not None and str(v).strip() != "" for v in r)]
    return name, headers, rows


def _read_csv(content: bytes) -> tuple[None, list[str], list[list]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DispatchEventImportError(f"Could not read CSV file as UTF-8: {exc}") from exc
    reader = csv.reader(io.StringIO(text))
    try:
        headers = [h.strip() for h in next(reader)]
    except StopIteration:
        raise DispatchEventImportError("File is empty")
    return None, headers, [row for row in reader if any(cell.strip() for cell in row)]


def _text(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _time(value, label: str) -> Optional[datetime]:
    """No time-zone conversion: naive times are taken as plant local time, like the telemetry uploads."""
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value
    try:
        return date_parser.parse(str(value).strip())
    except (ValueError, OverflowError):
        raise ValueError(f"{label} '{value}' is not a date/time")


def _number(value, label: str) -> Optional[Decimal]:
    if value is None or str(value).strip() == "":
        return None
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError(f"{label} '{value}' is not a number")
    if not number.is_finite():
        raise ValueError(f"{label} '{value}' is not a number")
    return number


def _bool(value, label: str) -> Optional[bool]:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        return value
    word = str(value).strip().lower()
    if word in TRUE_WORDS:
        return True
    if word in FALSE_WORDS:
        return False
    raise ValueError(f"{label} '{value}' is not TRUE or FALSE")


def _naive(t: Optional[datetime]) -> Optional[datetime]:
    return t.replace(tzinfo=None) if t else None


def _span(start: datetime, end: Optional[datetime]) -> str:
    return f"{start:%Y-%m-%d %H:%M}–{end:%H:%M}" if end else f"{start:%Y-%m-%d %H:%M}, still active"


def find_overlap(windows: list[tuple[str, datetime, Optional[datetime]]], start: datetime,
                 end: Optional[datetime]) -> Optional[tuple[str, datetime, Optional[datetime]]]:
    """First window (id, from, to) that overlaps [start, end); an open end runs on indefinitely."""
    start, end = _naive(start), _naive(end)
    for other in windows:
        o_start, o_end = _naive(other[1]), _naive(other[2])
        if (end is None or o_start < end) and (o_end is None or start < o_end):
            return other
    return None


def supersede_running(windows: list[tuple[str, datetime, Optional[datetime]]], start: datetime):
    """A new dispatch starting at `start` ends a running one (no end) that started before it.

    Returns (windows with that one closed at `start`, the closed window's (id, from) or None).
    """
    for i, (other_id, other_start, other_end) in enumerate(windows):
        if other_end is None and _naive(other_start) < _naive(start):
            return windows[:i] + [(other_id, other_start, start)] + windows[i + 1:], (other_id, other_start)
    return windows, None


def parse_dispatch_event_file(filename: str, content: bytes, db: Session):
    """Returns (sheet, rows ready for upsert, row errors, total rows, ignored columns, columns to update,
    stored running dispatches to close as {dispatch_id: (effective_from, new end)})."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        sheet, headers, rows = _read_xlsx(content)
    elif lower.endswith(".csv"):
        sheet, headers, rows = _read_csv(content)
    else:
        raise DispatchEventImportError("Unsupported file type: upload an Excel (.xlsx) or CSV file")
    if len(rows) > MAX_ROWS:
        raise DispatchEventImportError(f"Too many rows ({len(rows)}); the limit is {MAX_ROWS}")

    columns: dict[int, str] = {}
    ignored: list[str] = []
    for i, header in enumerate(headers):
        key = _norm(header)
        field = COLUMN_MAP.get(key)
        if field and field not in columns.values():
            columns[i] = field
        elif header and key not in KNOWN_UNSTORED:
            ignored.append(header)
    missing = [f for f in REQUIRED if f not in columns.values()]
    if missing:
        raise DispatchEventImportError(
            f"Missing required column(s): {', '.join(missing)}. Found: {', '.join(h for h in headers if h)}"
        )

    present = set(columns.values())
    update_cols = sorted({col for field in present for col in FIELD_COLUMNS.get(field, ())})

    ems_by_code = {e.ems_code.upper(): e for e in db.query(models.EmsController)}
    modes = {m.mode_code for m in db.query(models.EmsDispatchMode)}
    services = {s.service_code for s in db.query(models.EmsGridService)}
    statuses = {s.lower(): s for s in DISPATCH_STATUSES}
    Event = models.DispatchEvent
    existing = {d: (t, e, to) for d, t, e, to in db.query(Event.dispatch_id, Event.time, Event.ems_id, Event.effective_to)}

    valid: list[dict] = []
    errors: list[schemas.BulkDispatchEventRowError] = []
    seen: dict[str, int] = {}
    row_of: dict[str, int] = {}

    for offset, raw in enumerate(rows):
        row_no = offset + 2  # 1-based, after the header row
        v = {field: (raw[i] if i < len(raw) else None) for i, field in columns.items()}
        problems: list[str] = []
        dispatch_id = _text(v.get("dispatch_id"))

        def attempt(fn, *args):
            try:
                return fn(*args)
            except ValueError as exc:
                problems.append(str(exc))
                return None

        start = attempt(_time, v.get("effective_from"), "effective_from")
        end = attempt(_time, v.get("effective_to"), "effective_to")
        issued = attempt(_time, v.get("issued_at"), "issued_at")
        numbers = {f: attempt(_number, v.get(f), f) for f in (
            "p_setpoint_kw", "target_soc_pct", "soc_start_pct", "soc_end_pct", "ramp_rate_kw_s",
            "import_kwh", "export_kwh", "avg_price_mwh", "revenue_usd", "duration_min", "net_energy_kwh")}
        accepted = attempt(_bool, v.get("accepted"), "accepted")

        if not dispatch_id:
            problems.append("dispatch_id is required")
        elif dispatch_id in seen:
            problems.append(f"duplicate of row {seen[dispatch_id]} (dispatch_id {dispatch_id})")
        if _text(v.get("effective_from")) is None:  # a time that doesn't parse is reported above
            problems.append("effective_from is required")

        ems_code = _text(v.get("ems_code"))
        ems = ems_by_code.get(ems_code.upper()) if ems_code else None
        if not ems_code:
            problems.append("ems_code is required")
        elif ems is None:
            problems.append(f"EMS '{ems_code}' is not in master data (known: {', '.join(sorted(ems_by_code)) or 'none'})")

        command = (_text(v.get("command_code")) or "").upper()
        if not command:
            problems.append("command_code is required")
        elif command not in COMMAND_CODES:
            problems.append(f"unknown command_code '{command}' (expected {', '.join(COMMAND_CODES)})")
        mode = (_text(v.get("dispatch_mode")) or "").upper()
        if not mode:
            problems.append("dispatch_mode is required")
        elif mode not in modes:
            problems.append(f"unknown dispatch_mode '{mode}' (expected {', '.join(sorted(modes))})")
        service = (_text(v.get("grid_service")) or "NONE").upper()
        if service not in services:
            problems.append(f"unknown grid_service '{service}' (expected {', '.join(sorted(services))})")

        for f in ("target_soc_pct", "soc_start_pct", "soc_end_pct"):
            if numbers[f] is not None and not 0 <= numbers[f] <= 100:
                problems.append(f"{f} {numbers[f]} is outside 0-100")
        for f in ("import_kwh", "export_kwh", "ramp_rate_kw_s"):
            if numbers[f] is not None and numbers[f] < 0:
                problems.append(f"{f} {numbers[f]} can't be negative")

        issued = issued or start
        # no end-time column: an existing dispatch keeps its stored end (and so its status)
        if "effective_to" not in present and dispatch_id in existing:
            end = existing[dispatch_id][2]
        if start and end and end < start:
            problems.append("effective_to is before effective_from")
        if start and issued and _naive(issued) > _naive(start):
            problems.append("issued_at is after effective_from")

        status_text = _text(v.get("status"))
        status = statuses.get(status_text.lower()) if status_text else ("Active" if end is None else "Completed")
        if status is None:
            problems.append(f"unknown status '{status_text}' (expected {', '.join(DISPATCH_STATUSES)})")
        elif status == "Active" and end is not None:
            problems.append("status is Active but effective_to is set")
        elif status != "Active" and end is None and start is not None:
            problems.append(f"status is {status} but effective_to is empty")

        duration = numbers["duration_min"]
        if duration is not None and start and end and abs(float(duration) - (end - start).total_seconds() / 60) > 1:
            problems.append(f"duration_min {duration} doesn't match effective_from/to ({(end - start).total_seconds() / 60:.0f} min)")
        imp, exp, net = numbers["import_kwh"] or Decimal(0), numbers["export_kwh"] or Decimal(0), numbers["net_energy_kwh"]
        if net is not None and abs(imp - exp - net) > NET_TOLERANCE_KWH:
            problems.append(f"net_energy_kwh {net} doesn't match import - export ({imp - exp})")

        if dispatch_id and start and dispatch_id in existing and _naive(existing[dispatch_id][0]) != _naive(start):
            problems.append(f"dispatch_id {dispatch_id} already exists with effective_from {existing[dispatch_id][0]:%Y-%m-%d %H:%M}")

        if problems:
            errors.append(schemas.BulkDispatchEventRowError(row=row_no, dispatch_id=dispatch_id, errors=problems))
            continue
        seen[dispatch_id] = row_no
        row_of[dispatch_id] = row_no
        valid.append(dict(
            dispatch_id=dispatch_id,
            time=start,
            ems_id=ems.ems_id,
            battery_system_id=ems.battery_system_id,
            command_code=command,
            dispatch_mode=mode,
            grid_service=service,
            source=_text(v.get("source")),
            issued_at=issued,
            effective_to=end,
            p_setpoint_kw=numbers["p_setpoint_kw"] or Decimal(0),
            target_soc_pct=numbers["target_soc_pct"],
            soc_start_pct=numbers["soc_start_pct"],
            soc_end_pct=numbers["soc_end_pct"],
            ramp_rate_kw_s=numbers["ramp_rate_kw_s"],
            import_kwh=imp,
            export_kwh=exp,
            avg_price_mwh=numbers["avg_price_mwh"],
            revenue_usd=numbers["revenue_usd"],
            accepted=True if accepted is None else accepted,
            status=status,
            note=_text(v.get("note")),
        ))

    # one dispatch at a time per EMS: check each window against the stored ones (other ids) and earlier file rows
    in_file = {r["dispatch_id"] for r in valid}
    kept: list[dict] = []
    windows: dict[int, list] = {}
    for d, (t, ems_id, to) in existing.items():
        if d not in in_file:
            windows.setdefault(ems_id, []).append((d, t, to))
    closes: dict[str, tuple[datetime, datetime]] = {}
    by_id = {r["dispatch_id"]: r for r in valid}
    for r in sorted(valid, key=lambda r: _naive(r["time"])):
        ems_windows, closed = supersede_running(windows.get(r["ems_id"], []), r["time"])
        if closed:
            windows[r["ems_id"]] = ems_windows
            earlier = by_id.get(closed[0])
            if earlier is not None and earlier in kept:
                # a running dispatch earlier in the same file
                earlier["effective_to"], earlier["status"] = r["time"], "Completed"
                update_cols = sorted(set(update_cols) | {"effective_to", "status"})
            else:
                closes[closed[0]] = (closed[1], r["time"])
        clash = find_overlap(windows.get(r["ems_id"], []), r["time"], r["effective_to"])
        if clash:
            errors.append(schemas.BulkDispatchEventRowError(
                row=row_of[r["dispatch_id"]], dispatch_id=r["dispatch_id"],
                errors=[f"overlaps {clash[0]} ({_span(clash[1], clash[2])}) on the same EMS"
                        + (f"; end {clash[0]} first" if clash[2] is None else "")]))
            continue
        windows.setdefault(r["ems_id"], []).append((r["dispatch_id"], r["time"], r["effective_to"]))
        kept.append(r)
    errors.sort(key=lambda e: e.row)
    return sheet, kept, errors, len(rows), ignored, update_cols, closes


def upsert_dispatch_events(db: Session, rows: list[dict], update_cols: Optional[list[str]] = None) -> tuple[int, int]:
    """Insert new (dispatch_id, time) rows and update existing ones. Returns (inserted, updated).

    `update_cols` limits what an existing row gets updated with (the columns the file had); None updates all.
    """
    if not rows:
        return 0, 0
    table = models.DispatchEvent.__table__
    inserted = updated = 0
    for start in range(0, len(rows), 500):
        chunk = rows[start : start + 500]
        stmt = pg_insert(table).values(chunk)
        cols = [c for c in chunk[0] if c not in ("dispatch_id", "time") and (update_cols is None or c in update_cols)]
        stmt = stmt.on_conflict_do_update(
            index_elements=["dispatch_id", "time"],
            set_={**{c: getattr(stmt.excluded, c) for c in cols}, "updated_at": func.now()},
        ).returning(literal_column("(xmax = 0)").label("was_insert"))
        for row in db.execute(stmt):
            if row.was_insert:
                inserted += 1
            else:
                updated += 1
    return inserted, updated
