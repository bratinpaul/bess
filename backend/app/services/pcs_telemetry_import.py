import csv
import io
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Optional

import openpyxl
from dateutil import parser as date_parser
from pydantic import ValidationError
from sqlalchemy import literal_column
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

from app import models, schemas

MAX_ROWS = 200_000

REQUIRED_HEADERS = {"pcs_code"}
TIMESTAMP_HEADERS = ("ts", "timestamp")

# Maps a normalized (lowercased) source column header to a pcs_telemetry field.
COLUMN_MAP = {
    "mode_code": "mode_code",
    "grid_service": "grid_service",
    "ac_active_power_kw": "ac_active_power_kw",
    "ac_reactive_power_kvar": "ac_reactive_power_kvar",
    "ac_apparent_power_kva": "ac_apparent_power_kva",
    "power_factor": "power_factor",
    "ac_frequency_hz": "ac_frequency_hz",
    "ac_voltage_ab_v": "ac_voltage_ab_v",
    "ac_voltage_bc_v": "ac_voltage_bc_v",
    "ac_voltage_ca_v": "ac_voltage_ca_v",
    "ac_current_a_a": "ac_current_a_a",
    "ac_current_b_a": "ac_current_b_a",
    "ac_current_c_a": "ac_current_c_a",
    "thd_voltage_pct": "thd_voltage_pct",
    "thd_current_pct": "thd_current_pct",
    "dc_voltage_v": "dc_voltage_v",
    "dc_current_a": "dc_current_a",
    "dc_power_kw": "dc_power_kw",
    "efficiency_pct": "efficiency_pct",
    "igbt_temp_c": "igbt_temp_c",
    "heatsink_temp_c": "heatsink_temp_c",
    "inlet_coolant_temp_c": "inlet_coolant_temp_c",
    "ambient_temp_c": "ambient_temp_c",
    "avail_charge_kw": "avail_charge_kw",
    "avail_discharge_kw": "avail_discharge_kw",
    "ac_energy_import_kwh": "ac_energy_import_kwh",
    "ac_energy_export_kwh": "ac_energy_export_kwh",
    "breaker_state": "breaker_state",
    "alarm_level": "alarm_level",
}

DECIMAL_FIELDS = {
    "ac_active_power_kw",
    "ac_reactive_power_kvar",
    "ac_apparent_power_kva",
    "power_factor",
    "ac_frequency_hz",
    "ac_voltage_ab_v",
    "ac_voltage_bc_v",
    "ac_voltage_ca_v",
    "ac_current_a_a",
    "ac_current_b_a",
    "ac_current_c_a",
    "thd_voltage_pct",
    "thd_current_pct",
    "dc_voltage_v",
    "dc_current_a",
    "dc_power_kw",
    "efficiency_pct",
    "igbt_temp_c",
    "heatsink_temp_c",
    "inlet_coolant_temp_c",
    "ambient_temp_c",
    "avail_charge_kw",
    "avail_discharge_kw",
    "ac_energy_import_kwh",
    "ac_energy_export_kwh",
}


class PcsTelemetryImportError(Exception):
    """Raised for file-level problems: bad format, missing columns, empty file."""


def _read_xlsx(content: bytes) -> tuple[list[str], list[list]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - surfaced to the caller as a clean import error
        raise PcsTelemetryImportError(f"Could not read Excel file: {exc}") from exc

    sheet_name = next((n for n in wb.sheetnames if n.strip().lower() == "data"), None)
    if sheet_name is None:
        if len(wb.sheetnames) == 1:
            sheet_name = wb.sheetnames[0]
        else:
            raise PcsTelemetryImportError(
                f"Could not find a 'Data' worksheet. Available sheets: {', '.join(wb.sheetnames)}"
            )
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        raise PcsTelemetryImportError("Worksheet is empty")
    rows = [list(r) for r in rows_iter if any(v is not None for v in r)]
    return headers, rows


def _read_csv(content: bytes) -> tuple[list[str], list[list]]:
    try:
        text_content = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise PcsTelemetryImportError(f"Could not read CSV file as UTF-8: {exc}") from exc
    reader = csv.reader(io.StringIO(text_content))
    try:
        headers = [h.strip() for h in next(reader)]
    except StopIteration:
        raise PcsTelemetryImportError("File is empty")
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    return headers, rows


def _parse_timestamp(value) -> datetime:
    if isinstance(value, datetime):
        return value
    if value is None or str(value).strip() == "":
        raise ValueError("missing timestamp")
    return date_parser.parse(str(value).strip())


def _parse_decimal(value) -> Optional[Decimal]:
    if value is None:
        return None
    text_value = str(value).strip()
    if text_value == "":
        return None
    try:
        return Decimal(text_value)
    except InvalidOperation:
        raise ValueError(f"'{value}' is not a valid number")


def _parse_int(value) -> Optional[int]:
    if value is None:
        return None
    text_value = str(value).strip()
    if text_value == "":
        return None
    try:
        return int(float(text_value))
    except ValueError:
        raise ValueError(f"'{value}' is not a valid integer")


def _parse_text(value) -> Optional[str]:
    if value is None:
        return None
    text_value = str(value).strip()
    return text_value or None


def parse_pcs_telemetry_file(
    filename: str, content: bytes, db: Session
) -> tuple[list["schemas.PcsTelemetryCreate"], list["schemas.BulkPcsTelemetryRowError"], int]:
    """Parse an uploaded xlsx/csv file into validated pcs_telemetry rows.

    Returns (valid_rows, row_errors, total_data_rows). Raises PcsTelemetryImportError
    for file-level problems (wrong type, missing required columns, unreadable file).
    """
    lower_name = filename.lower()
    if lower_name.endswith((".xlsx", ".xls")):
        headers, raw_rows = _read_xlsx(content)
    elif lower_name.endswith(".csv"):
        headers, raw_rows = _read_csv(content)
    else:
        raise PcsTelemetryImportError("Unsupported file type. Please upload a .xlsx or .csv file.")

    if len(raw_rows) > MAX_ROWS:
        raise PcsTelemetryImportError(f"File has {len(raw_rows)} data rows, which exceeds the {MAX_ROWS} row limit.")

    normalized_headers = [h.strip().lower() for h in headers]
    header_index = {h: i for i, h in enumerate(normalized_headers) if h}

    missing_required = REQUIRED_HEADERS - set(header_index)
    if missing_required:
        raise PcsTelemetryImportError(f"Missing required column(s): {', '.join(sorted(missing_required))}")

    ts_header = next((h for h in TIMESTAMP_HEADERS if h in header_index), None)
    if ts_header is None:
        raise PcsTelemetryImportError("Missing required column: ts (or Timestamp)")

    field_index = {field: header_index[header] for header, field in COLUMN_MAP.items() if header in header_index}
    pcs_code_idx = header_index["pcs_code"]
    ts_idx = header_index[ts_header]

    pcs_ids = {p.pcs_code.strip().lower(): p.pcs_id for p in db.query(models.Pcs).all()}
    known_modes = {m.mode_code for m in db.query(models.PcsMode).all()}
    known_grid_services = {g.service_code for g in db.query(models.PcsGridService).all()}
    known_breaker_states = {b.state_code for b in db.query(models.PcsBreakerState).all()}

    valid_rows: list[schemas.PcsTelemetryCreate] = []
    row_errors: list[schemas.BulkPcsTelemetryRowError] = []

    for i, raw in enumerate(raw_rows):
        row_number = i + 2  # header occupies row 1
        errors: list[str] = []

        def cell(idx: int):
            return raw[idx] if idx < len(raw) else None

        pcs_code = _parse_text(cell(pcs_code_idx))
        pcs_id = None
        if not pcs_code:
            errors.append("pcs_code is required")
        else:
            pcs_id = pcs_ids.get(pcs_code.lower())
            if pcs_id is None:
                errors.append(f"Unknown pcs_code '{pcs_code}' (no matching PCS unit)")

        ts = None
        try:
            ts = _parse_timestamp(cell(ts_idx))
        except (ValueError, OverflowError) as exc:
            errors.append(f"Invalid timestamp: {exc}")

        field_values: dict = {}
        for field, idx in field_index.items():
            raw_value = cell(idx)
            try:
                if field == "alarm_level":
                    field_values[field] = _parse_int(raw_value)
                elif field in DECIMAL_FIELDS:
                    field_values[field] = _parse_decimal(raw_value)
                else:
                    field_values[field] = _parse_text(raw_value)
            except ValueError as exc:
                errors.append(f"{field}: {exc}")

        mode_code = field_values.get("mode_code")
        if mode_code and mode_code not in known_modes:
            errors.append(f"Unknown mode_code '{mode_code}'")
        grid_service = field_values.get("grid_service")
        if grid_service and grid_service not in known_grid_services:
            errors.append(f"Unknown grid_service '{grid_service}'")
        breaker_state = field_values.get("breaker_state")
        if breaker_state and breaker_state not in known_breaker_states:
            errors.append(f"Unknown breaker_state '{breaker_state}'")

        if errors:
            row_errors.append(schemas.BulkPcsTelemetryRowError(row=row_number, errors=errors))
            continue

        try:
            record = schemas.PcsTelemetryCreate(pcs_id=pcs_id, ts=ts, **field_values)
        except ValidationError as exc:
            messages = [f"{'.'.join(str(p) for p in e['loc'])}: {e['msg']}" for e in exc.errors()]
            row_errors.append(schemas.BulkPcsTelemetryRowError(row=row_number, errors=messages))
            continue

        valid_rows.append(record)

    return valid_rows, row_errors, len(raw_rows)


def upsert_pcs_telemetry_rows(db: Session, rows: list["schemas.PcsTelemetryCreate"]) -> tuple[int, int]:
    """Insert new (pcs_id, ts) rows and update rows that already exist.

    Existing rows are updated only in the columns the rows carry; the other columns keep their values.

    Returns (inserted_count, updated_count). Uses Postgres's `xmax = 0` trick on the
    RETURNING clause to tell inserts from updates, since comparing naive vs.
    timezone-aware datetimes ourselves would be unreliable.
    """
    if not rows:
        return 0, 0

    # De-duplicate within the file itself: keep the last occurrence of a given key.
    dedup: dict[tuple, dict] = {}
    for r in rows:
        data = r.model_dump()
        dedup[(data["pcs_id"], data["ts"])] = data
    all_rows = list(dedup.values())

    table = models.PcsTelemetry.__table__
    # A row that already exists (same pcs_id + ts) is updated only in the columns the rows were given - for a
    # bulk file, the columns it has - so values a partial file doesn't carry keep what is stored.
    # (With no such column, ts = ts keeps it a no-op update, so the row still counts as updated.)
    sent = set().union(*(r.model_fields_set for r in rows))
    update_cols = [c.name for c in table.columns if c.name not in ("pcs_id", "ts") and c.name in sent] or ["ts"]

    inserted = 0
    updated = 0
    chunk_size = 500
    for start in range(0, len(all_rows), chunk_size):
        chunk = all_rows[start : start + chunk_size]
        stmt = pg_insert(table).values(chunk)
        stmt = stmt.on_conflict_do_update(
            index_elements=["pcs_id", "ts"],
            set_={name: getattr(stmt.excluded, name) for name in update_cols},
        )
        stmt = stmt.returning(literal_column("(xmax = 0)").label("was_insert"))
        result = db.execute(stmt)
        for row in result:
            if row.was_insert:
                inserted += 1
            else:
                updated += 1

    return inserted, updated
