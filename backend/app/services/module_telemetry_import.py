"""Parse and upsert module telemetry uploads (xlsx or csv) into module_telemetry.

Expected columns (case-insensitive, extra columns ignored):
    module_code, ts (or timestamp), cell_v_avg_v, cell_v_min_v, cell_v_max_v,
    cell_temp_max_c, cell_temp_min_c, imbalance_mv (optional, checked), balancing_active

module_code is resolved to battery_module.module_id: an exact serial_number match first,
otherwise the module number (MOD-07 -> the one module whose serial ends in 7).
Timestamps are stored as given (no UTC conversion).
"""
import io
import re
from decimal import Decimal
from typing import Optional

import openpyxl
from pydantic import ValidationError
from sqlalchemy import literal_column
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

from app import models, schemas
from app.services.telemetry_import import (
    MAX_ROWS,
    TelemetryImportError,
    _parse_decimal,
    _parse_text,
    _parse_timestamp,
    _read_csv,
)

DECIMAL_FIELDS = ["cell_v_avg_v", "cell_v_min_v", "cell_v_max_v", "cell_temp_max_c", "cell_temp_min_c"]
TIMESTAMP_HEADERS = ("ts", "timestamp")
# the file's imbalance_mv must agree with (max - min) x 1000 to within this many mV
IMBALANCE_TOLERANCE_MV = Decimal("0.5")

TRUE_WORDS = {"true", "t", "yes", "y", "1", "on"}
FALSE_WORDS = {"false", "f", "no", "n", "0", "off"}


def _read_module_xlsx(content: bytes) -> tuple[list[str], list[list]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - surfaced to the caller as a clean import error
        raise TelemetryImportError("Could not read the Excel file. Please upload a valid .xlsx workbook.",
                                   log_message=f"xlsx not readable: {type(exc).__name__}: {exc}") from exc

    # "Data", or a data sheet with a suffix such as "Data (01 Aug sample)"
    names = [n for n in wb.sheetnames if n.strip().lower() == "data"]
    names = names or [n for n in wb.sheetnames if n.strip().lower().startswith("data")]
    if not names:
        if len(wb.sheetnames) != 1:
            raise TelemetryImportError(
                f"Could not find a 'Data' worksheet. Available sheets: {', '.join(wb.sheetnames)}"
            )
        names = wb.sheetnames
    rows_iter = wb[names[0]].iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        raise TelemetryImportError("Worksheet is empty")
    rows = [list(r) for r in rows_iter if any(v is not None for v in r)]
    return headers, rows


def _parse_bool(value) -> Optional[bool]:
    if value is None or isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text == "":
        return None
    if text in TRUE_WORDS:
        return True
    if text in FALSE_WORDS:
        return False
    raise ValueError(f"'{value}' is not true/false")


def _trailing_number(text: str) -> Optional[int]:
    m = re.search(r"(\d+)\s*$", text)
    return int(m.group(1)) if m else None


class ModuleResolver:
    """Maps a module_code from the file to a battery_module.module_id."""

    def __init__(self, db: Session):
        modules = db.query(models.BatteryModule.module_id, models.BatteryModule.serial_number).all()
        self.by_serial = {(s or "").strip().lower(): mid for mid, s in modules if s}
        self.by_number: dict[int, list[int]] = {}
        for mid, serial in modules:
            n = _trailing_number(serial or "")
            if n is not None:
                self.by_number.setdefault(n, []).append(mid)

    def resolve(self, code: str) -> tuple[Optional[int], Optional[str]]:
        exact = self.by_serial.get(code.lower())
        if exact is not None:
            return exact, None
        n = _trailing_number(code)
        if n is None:
            return None, f"Unknown module_code '{code}' (no matching battery module)"
        matches = self.by_number.get(n, [])
        if len(matches) == 1:
            return matches[0], None
        if not matches:
            return None, f"Unknown module_code '{code}' (no battery module numbered {n})"
        return None, f"Ambiguous module_code '{code}': {len(matches)} battery modules are numbered {n}"


def parse_module_telemetry_file(
    filename: str, content: bytes, db: Session
) -> tuple[list["schemas.ModuleTelemetryCreate"], list["schemas.BulkTelemetryRowError"], int]:
    """Returns (valid_rows, row_errors, total_data_rows); raises TelemetryImportError for file-level problems."""
    lower_name = filename.lower()
    if lower_name.endswith((".xlsx", ".xls")):
        headers, raw_rows = _read_module_xlsx(content)
    elif lower_name.endswith(".csv"):
        headers, raw_rows = _read_csv(content)
    else:
        raise TelemetryImportError("Unsupported file type. Please upload a .xlsx or .csv file.")

    if len(raw_rows) > MAX_ROWS:
        raise TelemetryImportError(f"File has {len(raw_rows)} data rows, which exceeds the {MAX_ROWS} row limit.")

    header_index = {h.strip().lower(): i for i, h in enumerate(headers) if h and h.strip()}
    ts_header = next((h for h in TIMESTAMP_HEADERS if h in header_index), None)
    missing = [name for name, ok in (("module_code", "module_code" in header_index), ("ts", ts_header)) if not ok]
    if missing:
        raise TelemetryImportError(f"Missing required column(s): {', '.join(missing)}")
    if not any(f in header_index for f in DECIMAL_FIELDS + ["balancing_active"]):
        raise TelemetryImportError(
            "No telemetry columns found. Expected some of: " + ", ".join(DECIMAL_FIELDS + ["balancing_active"])
        )

    code_idx, ts_idx = header_index["module_code"], header_index[ts_header]
    imbalance_idx = header_index.get("imbalance_mv")
    resolver = ModuleResolver(db)

    valid_rows: list[schemas.ModuleTelemetryCreate] = []
    row_errors: list[schemas.BulkTelemetryRowError] = []

    for i, raw in enumerate(raw_rows):
        row_number = i + 2  # header occupies row 1
        errors: list[str] = []

        def cell(idx: Optional[int]):
            return raw[idx] if idx is not None and idx < len(raw) else None

        module_id = None
        code = _parse_text(cell(code_idx))
        if not code:
            errors.append("module_code is required")
        else:
            module_id, problem = resolver.resolve(code)
            if problem:
                errors.append(problem)

        ts = None
        try:
            ts = _parse_timestamp(cell(ts_idx))
        except (ValueError, OverflowError) as exc:
            errors.append(f"Invalid ts: {exc}")

        values: dict = {}
        for field in DECIMAL_FIELDS:
            if field in header_index:
                try:
                    values[field] = _parse_decimal(cell(header_index[field]))
                except ValueError as exc:
                    errors.append(f"{field}: {exc}")
        if "balancing_active" in header_index:
            try:
                values["balancing_active"] = _parse_bool(cell(header_index["balancing_active"]))
            except ValueError as exc:
                errors.append(f"balancing_active: {exc}")

        # imbalance_mv is derived, not stored: it must agree with the cell voltage spread
        if imbalance_idx is not None and not errors:
            try:
                imbalance = _parse_decimal(cell(imbalance_idx))
            except ValueError as exc:
                errors.append(f"imbalance_mv: {exc}")
                imbalance = None
            lo, hi = values.get("cell_v_min_v"), values.get("cell_v_max_v")
            if imbalance is not None and lo is not None and hi is not None:
                expected = (hi - lo) * 1000
                if abs(expected - imbalance) > IMBALANCE_TOLERANCE_MV:
                    errors.append(
                        f"imbalance_mv {imbalance} does not match (cell_v_max_v - cell_v_min_v) x 1000 = {expected}"
                    )

        if errors:
            row_errors.append(schemas.BulkTelemetryRowError(row=row_number, errors=errors))
            continue

        try:
            valid_rows.append(schemas.ModuleTelemetryCreate(module_id=module_id, ts=ts, **values))
        except ValidationError as exc:
            messages = [
                (f"{'.'.join(str(p) for p in e['loc'])}: " if e["loc"] else "") + e["msg"].removeprefix("Value error, ")
                for e in exc.errors()
            ]
            row_errors.append(schemas.BulkTelemetryRowError(row=row_number, errors=messages))

    return valid_rows, row_errors, len(raw_rows)


def upsert_module_telemetry_rows(db: Session, rows: list["schemas.ModuleTelemetryCreate"]) -> tuple[int, int]:
    """Insert new (module_id, ts) rows and update existing ones. Returns (inserted, updated).

    Existing rows are updated only in the columns the rows carry; the other columns keep their values.
    """
    if not rows:
        return 0, 0

    # de-duplicate within the file itself: the last occurrence of a key wins
    dedup: dict[tuple, dict] = {}
    for r in rows:
        data = r.model_dump()
        dedup[(data["module_id"], data["ts"])] = data
    all_rows = list(dedup.values())

    table = models.ModuleTelemetry.__table__
    # A row that already exists (same module_id + ts) is updated only in the columns the rows were given - for a
    # bulk file, the columns it has - so values a partial file doesn't carry keep what is stored.
    # (With no such column, ts = ts keeps it a no-op update, so the row still counts as updated.)
    sent = set().union(*(r.model_fields_set for r in rows))
    update_cols = [c.name for c in table.columns if c.name not in ("module_id", "ts") and c.name in sent] or ["ts"]

    inserted = updated = 0
    for start in range(0, len(all_rows), 500):
        stmt = pg_insert(table).values(all_rows[start : start + 500])
        stmt = stmt.on_conflict_do_update(
            index_elements=["module_id", "ts"],
            set_={name: getattr(stmt.excluded, name) for name in update_cols},
        ).returning(literal_column("(xmax = 0)").label("was_insert"))
        for row in db.execute(stmt):
            if row.was_insert:
                inserted += 1
            else:
                updated += 1
    return inserted, updated
