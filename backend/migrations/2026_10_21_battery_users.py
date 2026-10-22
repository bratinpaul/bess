"""Login / RBAC schema battery_users, seeded from the RBAC matrix of BESS_Technical_Review_Recommendations.xlsx.

    python 2026_10_21_battery_users.py            # rehearsal: run, verify, ROLLBACK
    python 2026_10_21_battery_users.py --commit   # run for real (also writes the JWT keys + service account file)

role            the 7 platform roles (Tenant Viewer is tenant-scoped)
privilege       one row per resource:action held by any role (actions: read, write, delete, execute, approve)
role_privilege  the RBAC matrix (role x privilege)
app_user        fastapi-users user table (UUID id, email, argon2 hash, flags) + role, tenant, must_change_password, lockout
refresh_token   rotating refresh tokens (hash only), grouped in families for reuse detection
auth_audit      login / refresh / logout / failure events

Seeded users (password 'Password', must_change_password = true):
    admin@bess.local (Admin), asset.engineer@bess.local (Asset Engineer), operator@bess.local (Operator),
    data.scientist@bess.local (Data Scientist), tenant.viewer@bess.local (Tenant Viewer -> Tata Power Renewable Energy Ltd)
plus batch.service@bess.local (Data Engineer, random password in keys/service_account.json) for the loaders.

Re-running is safe: roles / privileges / matrix are re-synced from the workbook; existing users keep their passwords.
"""
import json
import secrets
import sys
import uuid
from pathlib import Path

import psycopg2
from openpyxl import load_workbook
from pwdlib import PasswordHash

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "BESS DESIGN" / "BESS_Technical_Review_Recommendations.xlsx"
KEYS = ROOT / "keys"
from bess_common.db_settings import psycopg2_dsn  # noqa: E402  (DATABASE_URL from env / backend/.env)

DSN = psycopg2_dsn()

ROLES = [  # code, name, tenant scoped, description
    ("ADMIN", "Admin", False, "Full access, user administration, approvals"),
    ("ASSET_ENGINEER", "Asset Engineer", False, "Maintains asset master data and fault records"),
    ("OPERATOR", "Operator", False, "Monitors the plant, logs faults and reviews diagnoses"),
    ("DATA_ENGINEER", "Data Engineer", False, "Loads telemetry and runs batch jobs"),
    ("DATA_SCIENTIST", "Data Scientist", False, "Trains and evaluates ML models"),
    ("ANALYST", "Analyst", False, "Read-only analysis and reports"),
    ("TENANT_VIEWER", "Tenant Viewer", True, "Read-only, limited to the sites of its own tenant"),
]
LETTER = {"R": "read", "W": "write", "D": "delete", "X": "execute", "A": "approve"}
USERS = [  # email, full name, role code, tenant name
    ("admin@bess.local", "Platform Admin", "ADMIN", None),
    ("asset.engineer@bess.local", "Asset Engineer", "ASSET_ENGINEER", None),
    ("operator@bess.local", "Plant Operator", "OPERATOR", None),
    ("data.scientist@bess.local", "Data Scientist", "DATA_SCIENTIST", None),
    ("tenant.viewer@bess.local", "Tenant Viewer (Tata Power)", "TENANT_VIEWER", "Tata Power Renewable Energy Limited"),
]
SERVICE = ("batch.service@bess.local", "Batch service account", "DATA_ENGINEER")

DDL = """
CREATE SCHEMA IF NOT EXISTS battery_users;

CREATE TABLE IF NOT EXISTS battery_users.role (
    role_id          smallint GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    role_code        text    NOT NULL UNIQUE,
    role_name        text    NOT NULL UNIQUE,
    is_tenant_scoped boolean NOT NULL DEFAULT false,
    description      text
);

CREATE TABLE IF NOT EXISTS battery_users.privilege (
    privilege_id  integer GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    resource      text NOT NULL,
    action        text NOT NULL CHECK (action IN ('read', 'write', 'delete', 'execute', 'approve')),
    code          text GENERATED ALWAYS AS (resource || ':' || action) STORED UNIQUE,
    module_label  text,
    UNIQUE (resource, action)
);

CREATE TABLE IF NOT EXISTS battery_users.role_privilege (
    role_id      smallint NOT NULL REFERENCES battery_users.role (role_id) ON DELETE CASCADE,
    privilege_id integer  NOT NULL REFERENCES battery_users.privilege (privilege_id) ON DELETE CASCADE,
    PRIMARY KEY (role_id, privilege_id)
);

CREATE TABLE IF NOT EXISTS battery_users.app_user (
    id                   uuid PRIMARY KEY,
    email                varchar(320) NOT NULL UNIQUE,
    hashed_password      varchar(1024) NOT NULL,
    is_active            boolean NOT NULL DEFAULT true,
    is_superuser         boolean NOT NULL DEFAULT false,
    is_verified          boolean NOT NULL DEFAULT true,
    full_name            text,
    role_id              smallint NOT NULL REFERENCES battery_users.role (role_id),
    tenant_id            bigint REFERENCES battery_master.tenant (tenant_id) ON DELETE SET NULL,
    is_service_account   boolean NOT NULL DEFAULT false,
    must_change_password boolean NOT NULL DEFAULT true,
    failed_logins        integer NOT NULL DEFAULT 0,
    locked_until         timestamptz,
    last_login_at        timestamptz,
    created_at           timestamptz NOT NULL DEFAULT now(),
    updated_at           timestamptz NOT NULL DEFAULT now()
);

CREATE TABLE IF NOT EXISTS battery_users.refresh_token (
    token_id    uuid PRIMARY KEY,
    user_id     uuid NOT NULL REFERENCES battery_users.app_user (id) ON DELETE CASCADE,
    family_id   uuid NOT NULL,
    token_hash  text NOT NULL UNIQUE,              -- sha256 of the opaque token; the token itself is never stored
    issued_at   timestamptz NOT NULL DEFAULT now(),
    expires_at  timestamptz NOT NULL,              -- end of the family's absolute lifetime
    used_at     timestamptz,                       -- set when rotated
    revoked_at  timestamptz,
    replaced_by uuid,
    user_agent  text,
    ip          text
);
CREATE INDEX IF NOT EXISTS ix_refresh_family ON battery_users.refresh_token (family_id);
CREATE INDEX IF NOT EXISTS ix_refresh_user ON battery_users.refresh_token (user_id);

CREATE TABLE IF NOT EXISTS battery_users.auth_audit (
    audit_id   bigint GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    ts         timestamptz NOT NULL DEFAULT now(),
    user_id    uuid,
    email      text,
    event      text NOT NULL,      -- LOGIN, LOGIN_FAILED, LOCKED, REFRESH, REFRESH_REUSE, LOGOUT
    success    boolean NOT NULL,
    ip         text,
    user_agent text,
    detail     jsonb
);
CREATE INDEX IF NOT EXISTS ix_auth_audit_ts ON battery_users.auth_audit (ts DESC);

CREATE OR REPLACE VIEW battery_users.v_role_permissions AS
SELECT r.role_code, r.role_name, p.resource, p.action, p.code, p.module_label
FROM battery_users.role r
JOIN battery_users.role_privilege rp ON rp.role_id = r.role_id
JOIN battery_users.privilege p ON p.privilege_id = rp.privilege_id;
"""


def read_matrix() -> tuple[dict[tuple[str, str], str], dict[str, set[tuple[str, str]]]]:
    """(resource, action) -> module label, and role name -> {(resource, action)} from the 'RBAC Matrix' sheet."""
    ws = load_workbook(WORKBOOK, read_only=True, data_only=True)["RBAC Matrix"]
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    header = rows[0]
    role_cols = {i: header[i] for i in range(2, len(header)) if header[i] in {r[1] for r in ROLES}}
    labels: dict[tuple[str, str], str] = {}
    grants: dict[str, set[tuple[str, str]]] = {name: set() for name in role_cols.values()}
    for row in rows[1:]:
        if not row or not row[0] or not row[1]:
            continue
        module = str(row[0]).replace("›", ">").replace("�", ">")
        resource = str(row[1]).split(",")[0].split(":")[0].strip()
        for i, role in role_cols.items():
            cell = (row[i] or "-").strip()
            for ch in cell:
                if ch in LETTER:
                    key = (resource, LETTER[ch])
                    grants[role].add(key)
                    labels.setdefault(key, module)
    return labels, grants


def write_keys() -> None:
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    KEYS.mkdir(exist_ok=True)
    priv, pub = KEYS / "jwt_private.pem", KEYS / "jwt_public.pem"
    if priv.exists() and pub.exists():
        print("JWT keys already present - kept")
        return
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    priv.write_bytes(key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                                       serialization.NoEncryption()))
    pub.write_bytes(key.public_key().public_bytes(serialization.Encoding.PEM,
                                                  serialization.PublicFormat.SubjectPublicKeyInfo))
    print("JWT RS256 key pair written to", KEYS)


def main(commit: bool) -> None:
    labels, grants = read_matrix()
    hasher = PasswordHash.recommended()  # argon2id
    conn = psycopg2.connect(**DSN)
    service_password = None
    try:
        with conn.cursor() as cur:
            cur.execute(DDL)
            for code, name, scoped, desc in ROLES:
                cur.execute("""INSERT INTO battery_users.role (role_code, role_name, is_tenant_scoped, description)
                               VALUES (%s, %s, %s, %s) ON CONFLICT (role_code) DO UPDATE SET role_name = EXCLUDED.role_name,
                               is_tenant_scoped = EXCLUDED.is_tenant_scoped, description = EXCLUDED.description""",
                            (code, name, scoped, desc))
            for (resource, action), label in sorted(labels.items()):
                cur.execute("""INSERT INTO battery_users.privilege (resource, action, module_label) VALUES (%s, %s, %s)
                               ON CONFLICT (resource, action) DO UPDATE SET module_label = EXCLUDED.module_label""",
                            (resource, action, label))
            # re-sync only the workbook's roles (roles defined elsewhere, e.g. BATCH_SERVICE, keep their grants)
            cur.execute("""DELETE FROM battery_users.role_privilege WHERE role_id IN
                           (SELECT role_id FROM battery_users.role WHERE role_name = ANY(%s))""", (list(grants),))
            for role_name, keys in grants.items():
                for resource, action in keys:
                    cur.execute("""INSERT INTO battery_users.role_privilege (role_id, privilege_id)
                                   SELECT r.role_id, p.privilege_id FROM battery_users.role r, battery_users.privilege p
                                   WHERE r.role_name = %s AND p.resource = %s AND p.action = %s""",
                                (role_name, resource, action))
            # privileges no role holds any more
            cur.execute("""DELETE FROM battery_users.privilege p WHERE NOT EXISTS
                           (SELECT 1 FROM battery_users.role_privilege rp WHERE rp.privilege_id = p.privilege_id)""")

            default_hash = hasher.hash("Password")
            for email, full_name, role_code, tenant in USERS:
                cur.execute("""INSERT INTO battery_users.app_user (id, email, hashed_password, full_name, role_id, tenant_id,
                                   is_superuser, must_change_password)
                               SELECT %s, %s, %s, %s, r.role_id,
                                      (SELECT tenant_id FROM battery_master.tenant WHERE name = %s), %s, true
                               FROM battery_users.role r WHERE r.role_code = %s
                               ON CONFLICT (email) DO UPDATE SET full_name = EXCLUDED.full_name, role_id = EXCLUDED.role_id,
                                   tenant_id = EXCLUDED.tenant_id""",
                            (str(uuid.uuid4()), email, default_hash, full_name, tenant, role_code == "ADMIN", role_code))
            cur.execute("SELECT 1 FROM battery_users.app_user WHERE email = %s", (SERVICE[0],))
            if cur.fetchone() is None:
                service_password = secrets.token_urlsafe(32)
                cur.execute("""INSERT INTO battery_users.app_user (id, email, hashed_password, full_name, role_id,
                                   is_service_account, must_change_password)
                               SELECT %s, %s, %s, %s, role_id, true, false FROM battery_users.role WHERE role_code = %s""",
                            (str(uuid.uuid4()), SERVICE[0], hasher.hash(service_password), SERVICE[1], SERVICE[2]))

            cur.execute("""SELECT r.role_name, count(rp.privilege_id) FROM battery_users.role r
                           LEFT JOIN battery_users.role_privilege rp ON rp.role_id = r.role_id GROUP BY r.role_id ORDER BY r.role_id""")
            print("privileges per role:", cur.fetchall())
            cur.execute("SELECT count(*) FROM battery_users.privilege")
            print("privileges:", cur.fetchone()[0])
            cur.execute("""SELECT u.email, r.role_name, t.name FROM battery_users.app_user u JOIN battery_users.role r USING (role_id)
                           LEFT JOIN battery_master.tenant t ON t.tenant_id = u.tenant_id ORDER BY u.created_at, u.email""")
            for row in cur.fetchall():
                print("  user", row)
        if commit:
            conn.commit()
            write_keys()
            if service_password:
                (KEYS / "service_account.json").write_text(json.dumps(
                    {"email": SERVICE[0], "password": service_password}, indent=2), encoding="utf-8")
                print("service account credentials written to", KEYS / "service_account.json")
            print("committed")
        else:
            conn.rollback()
            print("rehearsal - rolled back (use --commit)")
    finally:
        conn.close()


if __name__ == "__main__":
    main("--commit" in sys.argv)
